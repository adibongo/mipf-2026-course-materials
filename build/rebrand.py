import sys, re, shutil
from pptx import Presentation
from docx import Document
SRC,OUT=sys.argv[1],sys.argv[2]
REP=[("CIBN Centre for Financial Studies  |  Bankers House, Victoria Island","MIPF Boot Camp  |  Lagos Business School, Ajah"),
 ("CIBNCFS Workshop","MIPF Boot Camp, Day 1"),("23 September 2026","14 October 2026"),("23 Sept 2026","14 Oct 2026"),
 ("Wednesday, 23 September 2026","Wednesday, 14 October 2026"),
 ("LAGOS BUSINESS SCHOOL  |  CIBN CENTRE FOR FINANCIAL STUDIES","LAGOS BUSINESS SCHOOL  |  MIPF BOOT CAMP, DAY 1"),
 ("Workshop: Infrastructure Finance:","Day 1: Infrastructure Finance:"),
 ("Bankers House, Victoria Island, Lagos","Lagos Business School, Ajah, Lagos"),
 ("For the CIBNCFS workshop, Modules 1 & 2","For MIPF Day 1, Modules 1 & 2"),
 ("before the workshop","before Day 1"),("The workshop builds","Day 1 builds"),
 ("These map one-to-one to CIBN's objectives for Modules 1 and 2.","These are the learning objectives for Modules 1 and 2."),
 ("CIBN's last objective","The last objective"),("Thank CIBNCFS and participants.","Thank participants."),
 ("(9:10–9:50, 40 min)","(9:10–10:00)"),("(9:50–10:35, 45 min)","(10:00–11:00)"),
 ("MODULE 1 OPENER (9:10)","MODULE 1 OPENER (9:00)")]
def fix_para_runs(runs):
    changed=0
    joined=''.join(r.text for r in runs)
    if 'badi@lbs.edu.ng' in joined and '8699' not in joined:
        for r in runs:
            if 'badi@lbs.edu.ng' in r.text: r.text=r.text.replace('badi@lbs.edu.ng','badi@lbs.edu.ng  |  +234 706 164 8699  |  bongo.adi@fulbrightmail.org'); changed=1; break
    for r in runs:
        t=r.text
        for a,b in REP:
            if a in t: t=t.replace(a,b)
        if t!=r.text: r.text=t; changed=1
    full="".join(r.text for r in runs)
    # spanning phrases
    nf=full
    for a,b in REP: nf=nf.replace(a,b)
    if nf!=full and runs:
        runs[0].text=nf
        for r in runs[1:]: r.text=""
        changed=1
    return changed
def set_cell(c,txt):
    runs=[r for p in c.text_frame.paragraphs for r in p.runs] if hasattr(c,"text_frame") else [r for p in c.paragraphs for r in p.runs]
    if runs:
        runs[0].text=txt
        for r in runs[1:]: r.text=""
# ---- deck
p=Presentation(SRC+"/CIBN_Infrastructure_Finance_Modules_1-2_Deck.pptx")
AG=[("9:00 – 9:10",),("9:10 – 10:00",),("10:00 – 11:00",),("11:00 – 11:30",),("11:30 – 12:30",),("12:30 – 2:00",),("2:00 – 2:45",),("2:45 – 3:30",),("3:30 – 4:30",)]
def walk(shapes):
    for sh in shapes:
        if sh.shape_type==6: walk(sh.shapes)
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs: fix_para_runs(para.runs)
        if getattr(sh,"has_table",False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    for para in c.text_frame.paragraphs: fix_para_runs(para.runs)
for i,s in enumerate(p.slides,1):
    walk(s.shapes)
    if s.has_notes_slide:
        for para in s.notes_slide.notes_text_frame.paragraphs: fix_para_runs(para.runs)
    if i==2:
        for sh in s.shapes:
            if getattr(sh,"has_table",False) and sh.has_table:
                rows=list(sh.table.rows)[1:]
                for r,(t,) in zip(rows,AG): set_cell(r.cells[0],t)
                set_cell(rows[7].cells[1],"M2.2  PPPs & blended finance; M2.3  ESG")
                set_cell(rows[7].cells[2],"2nd Niger Bridge; tax credits; IFC PS, Equator")
                set_cell(rows[8].cells[1],"M2.4  Negotiation & credit committee")
                set_cell(rows[3].cells[1],"Coffee break")
    if i==7:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip()=="9:10 – 11:30": set_cell(sh,"9:10 – 12:30")
from pptx.util import Inches
LG=OUT+"/website/assets/logo/"
for s_ in p.slides:
    s_.shapes.add_picture(LG+"mipf_hex_256.png",Inches(8.55),Inches(0.1),height=Inches(0.42))
    s_.shapes.add_picture(LG+"bongo_ai_hex_256.png",Inches(9.05),Inches(0.1),height=Inches(0.42))
p.save(OUT+"/website/materials/MIPF_Day1_Modules_1-2_Deck.pptx")
# ---- docx
def do_docx(src,dst,table_times=None):
    d=Document(src)
    ps=list(d.paragraphs)+[q for t in d.tables for r in t.rows for c in r.cells for q in c.paragraphs]
    for sct in d.sections: ps+=list(sct.header.paragraphs)+list(sct.footer.paragraphs)
    for q in ps: fix_para_runs(q.runs)
    if table_times:
        t=d.tables[0]
        for r,new in zip(list(t.rows)[1:],table_times):
            c=r.cells[0]; rs=[x for q in c.paragraphs for x in q.runs]
            rs[0].text=new
            for x in rs[1:]: x.text=""
    from docx.shared import Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    LG=OUT+"/website/assets/logo/"
    hdr=d.sections[0].header; hp=hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
    hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    hp.add_run().add_picture(LG+"mipf_hex_256.png",height=Cm(0.9)); hp.add_run("  "); hp.add_run().add_picture(LG+"bongo_ai_hex_256.png",height=Cm(0.9))
    d.save(dst)
do_docx(SRC+"/CIBN_Participant_Workbook.docx",OUT+"/website/materials/MIPF_Day1_Participant_Workbook.docx",
        ["9:00 – 10:00","10:00 – 11:00","11:30 – 12:30","12:30 – 2:00","2:45 – 3:30","2:00 – 2:45","3:30 – 4:30"])
do_docx(SRC+"/CIBN_Pre-read_Primer.docx",OUT+"/website/materials/MIPF_Day1_PreRead_Primer.docx")
shutil.copy(SRC+"/Sahel_Sun_Capital_Stack_Template.xlsx",OUT+"/website/materials/")
shutil.copy("/Users/adibongo/Library/CloudStorage/OneDrive-Personal/Mastering Infrastructure Project Finance 2023/TrainStat_StartModel - Copy.xlsx",OUT+"/website/materials/TrainStat_StartModel.xlsx")
print("ok")
