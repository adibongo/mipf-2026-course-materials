# -*- coding: utf-8 -*-
"""Workbook surgery helpers: row insertion with global reference shifting, row fill, style copy."""
import re, copy
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.formatting.formatting import ConditionalFormattingList

FIRST, LAST = 10, 93   # columns J..CO

_part = re.compile(r"^(\$?[A-Za-z]{1,3})(\$?)(\d+)$")

def _shift_ref(ref, at, n):
    parts = ref.split(":")
    out = []
    for p in parts:
        m = _part.match(p)
        if not m: return ref
        row = int(m.group(3))
        if row >= at: row += n
        out.append(f"{m.group(1)}{m.group(2)}{row}")
    return ":".join(out)

def _shift_formula(f, own_sheet, target, at, n):
    if not isinstance(f, str) or not f.startswith("="): return f
    try:
        tok = Tokenizer(f)
    except Exception:
        return f
    out = []
    for t in tok.items:
        v = t.value
        if t.type == "OPERAND" and t.subtype == "RANGE":
            if "!" in v:
                sh, ref = v.rsplit("!", 1)
                shn = sh.strip("'")
                if shn == target: v = sh + "!" + _shift_ref(ref, at, n)
            elif own_sheet == target:
                v = _shift_ref(v, at, n)
        out.append(v)
    return "=" + "".join(out)

def insert_rows(wb, sheet, at, n=1, style_from=None):
    ws = wb[sheet]
    maxc = ws.max_column
    # row heights
    dims = {r: copy.copy(d) for r, d in ws.row_dimensions.items()}
    ws.insert_rows(at, n)
    for r in sorted([r for r in dims if r >= at], reverse=True):
        ws.row_dimensions[r + n].height = dims[r].height
    # conditional formatting
    old = ws.conditional_formatting
    new = ConditionalFormattingList()
    for rng in old:
        sq = str(rng.sqref)
        pieces = []
        for piece in sq.split():
            pieces.append(_shift_ref(piece, at, n))
        for rule in rng.rules:
            new.add(" ".join(pieces), rule)
    ws.conditional_formatting = new
    # global formula shifting
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula):
                    v.text = _shift_formula(v.text, w.title, sheet, at, n)
                    if w.title == sheet: v.ref = _shift_ref(v.ref, at, n)
                elif isinstance(v, str) and v.startswith("="):
                    c.value = _shift_formula(v, w.title, sheet, at, n)
    # defined names
    for name, dn in wb.defined_names.items():
        t = dn.attr_text
        if t and "!" in t:
            sh, ref = t.rsplit("!", 1)
            if sh.strip("'") == sheet:
                dn.attr_text = sh + "!" + _shift_ref(ref, at, n)
    if style_from is not None:
        sr = style_from if style_from < at else style_from + n
        for k in range(n):
            copy_row_style(ws, sr, at + k, maxc)
    return ws

def copy_row_style(ws, src, dst, maxc=97):
    for c in range(1, maxc + 1):
        s = ws.cell(src, c); d = ws.cell(dst, c)
        if s.has_style: d._style = copy.copy(s._style)
    if ws.row_dimensions[src].height: ws.row_dimensions[dst].height = ws.row_dimensions[src].height

def fill_row(ws, r, tmpl, first=FIRST, last=LAST):
    """tmpl is the formula for column J (or `first`); translated across."""
    org = f"{L(first)}{r}"
    for c in range(first, last + 1):
        ws.cell(r, c).value = tmpl if c == first else Translator(tmpl, origin=org).translate_formula(f"{L(c)}{r}")

def pull_row(ws, r, sheet_src, src_row, timeline=False, d=None, style_from=None):
    """Label/constant/unit pull (E:G) and optionally the timeline (J:CO) from sheet_src row src_row."""
    if style_from: copy_row_style(ws, style_from, r)
    if d is not None: ws.cell(r, 4).value = d
    for col in "EFG":
        ws[f"{col}{r}"].value = f"={sheet_src}!{col}${src_row}"
    if timeline:
        fill_row(ws, r, f"={sheet_src}!J${src_row}")
        ws[f"H{r}"].value = f"={sheet_src}!H${src_row}"
