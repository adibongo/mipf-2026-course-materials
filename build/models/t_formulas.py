import openpyxl, formulas, warnings, time
warnings.filterwarnings("ignore")
wb=openpyxl.Workbook(); ws=wb.active; ws.title="T"
ws["A1"]=100; ws["A2"]=0.05
for i in range(3,13): ws.cell(i,1).value=f"=A{i-1}*(1+$A$2)"
ws["B1"]="=NPV(A2,A3:A12)"; ws["B2"]="=IF(AND(A1>0,A2<1),1,0)"; ws["B3"]="=SUMIFS(A3:A12,A3:A12,\">110\")"
ws["B4"]="=INDEX(A3:A12,MATCH(0.0,A3:A12,0))" 
ws["B5"]="=EOMONTH(DATE(2021,12,31),3)"; ws["B6"]="=MAX(MIN(A1,200),0)"; ws["B7"]="=SUMPRODUCT(A3:A12,A3:A12)"
ws["B8"]="=XIRR(C1:C3,D1:D3)"; ws["C1"]=-100; ws["C2"]=60; ws["C3"]=70
import datetime
ws["D1"]=datetime.date(2021,12,31); ws["D2"]=datetime.date(2022,12,31); ws["D3"]=datetime.date(2023,12,31)
wb.save("t.xlsx")
t=time.time()
xl=formulas.ExcelModel().loads("t.xlsx").finish(); sol=xl.calculate()
for k,v in sol.items():
    if k.startswith("'[t.xlsx]T'!B"): print(k, v.value[0,0] if hasattr(v,'value') else v)
print(time.time()-t)
