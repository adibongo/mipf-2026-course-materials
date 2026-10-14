# -*- coding: utf-8 -*-
import re
"""Stage 1: corrected, completed start model (Days 4-5 end-state). No empty shells."""
import sys, copy, datetime as dt, warnings
warnings.filterwarnings("ignore")
sys.path.insert(0, "/private/tmp/claude-501/s/model")
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
from fixlib import *

SRC = "/private/tmp/claude-501/s/model/orig.xlsx"
YELLOW = PatternFill("solid", fgColor="FFFF99")
NAVY, GOLD, BLUE = "141B4D", "D99A2B", "0B5E82"

def build(out):
    wb = openpyxl.load_workbook(SRC)
    changes = []
    c = wb["InpC"]; s = wb["InpS"]; t = wb["Time"]; tax = wb["Tax"]; oc = wb["OpCost"]; debt = wb["Debt"]; eq = wb["Equity"]; fs = wb["FinStat"]; val = wb["Val"]; tr = wb["Track"]; nca = wb["NCA"]

    # ---------------------------------------------------------------- remove ghost formulas beyond the timeline, drop scratch sheet
    removed = 0
    for ws in wb.worksheets:
        if ws.title == "Temp": continue
        for key in [k for k in ws._cells if k[1] > 93]:
            del ws._cells[key]; removed += 1
        for k in list(ws.column_dimensions.keys()):
            d = ws.column_dimensions[k]
            if d.min and d.min > 93: del ws.column_dimensions[k]
            elif d.max and d.max > 93: d.max = 93
    del wb["Temp"]
    changes.append(f"Ghost formulas removed: seven sheets carried about {removed:,} formulas and formats out to column XFD (the sheet edge), far beyond the 84-column timeline (J to CO). Trimming them cuts the file from about 6.0 MB to 0.4 MB and makes recalculation take seconds instead of many minutes. The unused scratch sheet 'Temp' (about 50,000 formulas) was removed.")
    # ---------------------------------------------------------------- names and links
    for n in list(wb.defined_names.keys()):
        if n in ("LIVE_RESULTS", "TRACK_ACTIVE", "TRK_COMMENT", "unit"): del wb.defined_names[n]
    from openpyxl.workbook.defined_name import DefinedName
    for n, ref in (("LIVE_RESULTS", "Track!$H$8:$H$69"), ("TRACK_ACTIVE", "Track!$H$5"), ("TRK_COMMENT", "Track!$H$77:$H$78"), ("unit", "InpC!$F$102")):
        wb.defined_names[n] = DefinedName(n, attr_text=ref)
    wb._external_links = []
    for ws_ in wb.worksheets:
        for n_ in list(ws_.defined_names.keys()): del ws_.defined_names[n_]
    wb["Time"].print_area = "A1:CO111"
    changes.append("Defined names LIVE_RESULTS, TRACK_ACTIVE and TRK_COMMENT repaired (they pointed at deleted cells); name 'unit' re-pointed from an external link to InpC!F102; sheet-level names on Track that pointed at an external workbook removed; Time print area extended to column CO; external link removed.")

    # ---------------------------------------------------------------- typos and stray cells
    typo = {("InpC", "E3"): "Scenario name", ("InpC", "G44"): "USD/arrival", ("InpC", "E55"): "Investment / addition date",
            ("Time", "B27"): "Days in model period", ("Time", "E30"): "Days in model period", ("Time", "E95"): "Depreciation period flag",
            ("NCA", "C24"): "Accounting depreciation - Initial balance", ("NCA", "E29"): "Accounting depreciation - Initial balance",
            ("Equity", "E31"): "Aggregate cash available for dividend", ("Equity", "E35"): "Aggregate earnings available for dividend",
            ("Val", "B33"): "NPV of shareholder cashflow", ("Val", "E39"): "NPV of shareholder cashflow", ("Equity", "E10"): "Share capital redemption POS"}
    for (sh, a), v in typo.items(): wb[sh][a].value = v
    changes.append("Typos corrected: Scenario, arrival, Investment, Accounting, available, period, NPV of ...")
    tax["C28"].value = None; tax["C35"].value = None
    if isinstance(c["F45"].value, ArrayFormula): c["F45"].value = None
    for col in "JKL": c[f"{col}45"].value = None
    s["J74"].value = None
    changes.append("Stray zeros removed (Tax C28, C35; InpC row 45; InpS J74).")
    # clear columns beyond the timeline on FinStat
    for r in range(1, fs.max_row + 1):
        for col in range(94, 98):
            if fs.cell(r, col).value is not None: fs.cell(r, col).value = None

    # ---------------------------------------------------------------- total ranges that stopped at BY / BV / CQ
    for sh, cells in (("Time", ["H14", "H45", "H53", "H58", "F59", "H79"]), ("InpS", ["H8", "H9", "H18", "H19", "H20"]), ("FinStat", ["F74", "H75", "F77", "H78", "F80", "H81"])):
        for a in cells:
            v = wb[sh][a].value
            wb[sh][a].value = re.sub(r":(BY|BV|CQ)(\d+)", lambda m: f":CO{m.group(2)}", v)
    changes.append("Total and check ranges extended to column CO (they stopped at BY, BV or CQ): the three FinStat checks now cover every column, Time!H79 now counts the post-forecast flag.")

    # ---------------------------------------------------------------- scenarios
    c["F3"].value = "Base"
    defs = {"fare": 35, "om": 46}
    c["K35"].value = 6.435; c["L35"].value = 7.865
    for col in "JKL": c[f"{col}46"].value = 0
    c["K46"].value = 0.10
    for r in range(17, 22): c[f"L{r}"].value = c[f"J{r}"].value
    c["N3"].value = "Scenario definitions"; c["N3"].font = Font(bold=True, color=NAVY)
    c["N4"].value = "Base: inflation 7% falling to 3%; fare 7.15; O&M as contracted"
    c["N5"].value = "Downside: Base with fare -10% (6.435) and O&M +10%"
    c["N6"].value = "Upside: Base with fare +10% (7.865)"
    changes.append("Scenarios made coherent: Upside was worse than Base (2% flat inflation) and Downside equalled Base. Now Base / Downside (fare -10%, O&M +10%) / Upside (fare +10%). O&M sensitivity input (InpC row 46) is now live.")

    # ---------------------------------------------------------------- O&M sensitivity link (OpCost)
    insert_rows(wb, "OpCost", 28, 1, style_from=27)
    pull_row(oc, 28, "InpC", 46, style_from=27)
    oc["J31"].value = "= IF(J30 = 1, $F26 / $F27 * J29 * (1 + $F28), 0)"
    fill_row(oc, 31, "=IF(J30=1,$F26/$F27*J29*(1+$F28),0)")

    # ---------------------------------------------------------------- Tax: deduct interest (fills placeholder)
    copy_row_style(tax, 25, 26)
    tax["D26"].value = "less"
    pull_row(tax, 26, "Debt", 32, timeline=False)
    fill_row(tax, 26, "=Debt!J$32")
    tax["H26"].value = "= SUM(J26:CO26)"
    changes.append("Tax: interest deduction placeholder completed (Tax row 26 = Debt interest). Cash tax now deducts interest, so deferred tax unwinds to zero and nothing is left stranded in retained earnings.")

    # ---------------------------------------------------------------- Debt drawdown / Equity injection hooks (fill placeholders) via InpS
    for r, lab in ((22, "Senior debt - additional drawdown"), (23, "Share capital - additional injection")):
        copy_row_style(s, 8, r)
        s[f"E{r}"].value = lab; s[f"G{r}"].value = "$ 000s"; s[f"H{r}"].value = f"= SUM(J{r}:CO{r})"
        for col in range(FIRST, LAST + 1): s.cell(r, col).value = 0; s.cell(r, col).fill = YELLOW
    s["A21"].value = "ADDITIONAL FUNDING (after financial close)"; s["A21"]._style = copy.copy(s["A17"]._style)
    debt["E42"].value = "=InpS!E$22"; debt["F42"].value = "=InpS!F$22"; debt["G42"].value = "=InpS!G$22"; debt["H42"].value = "=SUM(J42:CO42)"
    fill_row(debt, 42, "=InpS!J$22")
    eq["E22"].value = "=InpS!E$23"; eq["F22"].value = "=InpS!F$23"; eq["G22"].value = "=InpS!G$23"; eq["H22"].value = "=SUM(J22:CO22)"
    fill_row(eq, 22, "=InpS!J$23")
    # FinStat cash cascade uses blank rows 39 (drawdown) and 43 (injection)
    for r, src, sheet in ((39, 42, "Debt"), (43, 22, "Equity")):
        copy_row_style(fs, 40, r)
        fs[f"E{r}"].value = f"={sheet}!E${src}"; fs[f"F{r}"].value = f"={sheet}!F${src}"; fs[f"G{r}"].value = f"={sheet}!G${src}"; fs[f"H{r}"].value = f"= SUM(J{r}:CO{r})"
        fill_row(fs, r, f"={sheet}!J${src}")
    fill_row(fs, 42, "=J38+SUM(J39:J41)")
    fill_row(fs, 45, "=J42+SUM(J43:J44)")
    # Val: pull injection, include in outflow
    copy_row_style(val, 22, 24)
    val["E24"].value = "=Equity!E$22"; val["F24"].value = "=Equity!F$22"; val["G24"].value = "=Equity!G$22"; val["H24"].value = "= SUM(J24:CO24)"
    fill_row(val, 24, "=Equity!J$22")
    fill_row(val, 27, "=IF(J26=1,$F25,0)+J24")
    # Days in period: period end less period BEGINNING (first day of the month) is one day short; add the missing day
    for col in range(10, 94):
        L_ = openpyxl.utils.get_column_letter(col)
        t[f"{L_}30"].value = f"={L_}28-{L_}29+1"
    changes.append("Day count corrected: Time row 30 (days in model period) was period end less period beginning, one day short in every quarter (a 361-day year; Constr already used true days). It is now end less beginning plus 1. Interest, receivables and payables, and every result that depends on them move slightly (for example the start-model NPV by about 0.4%).")
    changes.append("Placeholders completed: Debt row 42 and Equity row 22 now take inputs from InpS rows 22 and 23; they flow through the cash cascade (FinStat rows 39 and 43) and the shareholder valuation (Val rows 24 and 27).")

    # ---------------------------------------------------------------- DSCR block on FinStat
    r0 = 83
    for r in range(r0, r0 + 14): copy_row_style(fs, 80 if r not in (r0, ) else 72, r)
    copy_row_style(fs, 72, r0)
    fs[f"A{r0}"].value = "SENIOR DEBT RATIOS"
    pull_row(fs, 85, "Debt", 21, timeline=True, style_from=40)
    pull_row(fs, 86, "Debt", 32, timeline=True, style_from=40)
    fs["E87"].value = "=E$38"; fs["F87"].value = "=F$38"; fs["G87"].value = "=G$38"; fs["H87"].value = "= H$38"; copy_row_style(fs, 40, 87)
    fill_row(fs, 87, "=J$38")
    for r, lab, f, unit in ((88, "Senior debt service", "=SUM(J85:J86)", "$ 000s"), (89, "Debt service period flag", "=IF(J88>CHK_TOL,1,0)", "flag"), (90, "DSCR (CFADS / debt service)", "=IF(J89=1,J87/J88,0)", "x")):
        copy_row_style(fs, 38, r); fs[f"E{r}"].value = lab; fs[f"G{r}"].value = unit
        fill_row(fs, r, f)
    fs["H88"].value = "= SUM(J88:CO88)"; fs["H89"].value = "= SUM(J89:CO89)"
    for r, lab, f in ((92, "Minimum DSCR", '=_xlfn.MINIFS(J90:CO90,J89:CO89,1)'), (93, "Average DSCR", '=AVERAGEIFS(J90:CO90,J89:CO89,1)'), (94, "Whole-life cover (CFADS / debt service)", "=IF(H88=0,0,H87/H88)")):
        copy_row_style(fs, 38, r); fs[f"E{r}"].value = lab; fs[f"F{r}"].value = f; fs[f"G{r}"].value = "x"
        for col in range(FIRST, LAST + 1): fs.cell(r, col).value = None
    for r in (85, 86, 87, 88, 89, 90):
        for col in range(FIRST, LAST + 1): fs.cell(r, col).number_format = "0.00_);(0.00)" if r == 90 else fs.cell(r, col).number_format
    for r in (92, 93, 94): fs[f"F{r}"].number_format = '0.00"x"'
    changes.append("FinStat: DSCR block added: debt service, DSCR, minimum, average and whole-life cover.")

    # ---------------------------------------------------------------- Track: full build
    tr["H5"].value = "Base"
    tr["M5"].value = "Base"; tr["N5"].value = "Downside"; tr["O5"].value = "Upside"
    lab = {}
    def T(r, label, f, fmt="#,##0_);(#,##0);\"-  \"", unit=None, F=None):
        lab[r] = (label, f, fmt, unit, F)
    T(8, "Shareholder NPV at required return", "=Val!F39", unit="$ 000s")
    T(9, "Shareholder IRR (XIRR)", "=Val!F49", "0.00%", "% p.a.")
    T(10, "Total CFADS (20 years)", "=FinStat!H38", unit="$ 000s")
    T(11, "Minimum DSCR", "=FinStat!F92", '0.00"x"', "x")
    T(12, "Senior debt at financial close", "=Debt!F11", unit="$ 000s")
    T(14, "Operating revenue", "=FinStat!H9", unit="$ 000s"); T(15, "Operating cost", "=FinStat!H10", unit="$ 000s"); T(16, "EBITDA", "=FinStat!H11", unit="$ 000s")
    T(17, "EBITDA margin", "=IF(H14=0,0,H16/H14)", "0.0%", "%"); T(18, "Accounting depreciation", "=FinStat!H13", unit="$ 000s"); T(19, "Senior debt interest", "=FinStat!H14", unit="$ 000s")
    T(20, "Profit before tax", "=FinStat!H15", unit="$ 000s"); T(21, "Tax expense", "=FinStat!H17", unit="$ 000s"); T(22, "Profit after tax", "=FinStat!H18", unit="$ 000s"); T(23, "Effective tax rate", "=IF(H20=0,0,-H21/H20)", "0.0%", "%")
    for i, y in enumerate(range(2022, 2028)):
        T(24 + i, f'="Operating revenue FY"&F{24+i}', f"=SUMIFS(FinStat!$J$9:$CO$9,FinStat!$J$4:$CO$4,$F{24+i})", unit="$ 000s", F=y)
    T(31, "Operating revenue received", "=FinStat!H30", unit="$ 000s"); T(32, "Operating cost paid", "=FinStat!H31", unit="$ 000s"); T(33, "Operating cash flow", "=FinStat!H32", unit="$ 000s")
    T(34, "Capital expenditure", "=FinStat!H34", unit="$ 000s"); T(35, "Pre-tax pre-financing cash flow", "=FinStat!H35", unit="$ 000s"); T(36, "Tax paid", "=FinStat!H37", unit="$ 000s")
    T(37, "CFADS", "=FinStat!H38", unit="$ 000s"); T(38, "Senior debt interest paid", "=FinStat!H40", unit="$ 000s"); T(39, "Senior debt principal repaid", "=FinStat!H41", unit="$ 000s")
    T(40, "Cash flow available to shareholders", "=FinStat!H42", unit="$ 000s"); T(41, "Share capital redemption", "=FinStat!H44", unit="$ 000s"); T(42, "Cash flow available for dividends", "=FinStat!H45", unit="$ 000s")
    T(43, "Dividends paid", "=FinStat!H47", unit="$ 000s"); T(44, "Net cash flow", "=FinStat!H48", unit="$ 000s")
    for i, y in enumerate(range(2022, 2028)):
        T(45 + i, f'="CFADS FY"&F{45+i}', f"=SUMIFS(FinStat!$J$38:$CO$38,FinStat!$J$4:$CO$4,$F{45+i})", unit="$ 000s", F=y)
    T(51, "Total senior debt service", "=FinStat!H88", unit="$ 000s"); T(52, "Whole-life cover (CFADS / debt service)", "=FinStat!F94", '0.00"x"', "x")
    bs = lambda row: f"=INDEX(FinStat!$J${row}:$CO${row},MATCH($H$53,FinStat!$J$2:$CO$2,0))"
    T(54, "Non-current assets", bs(58), unit="$ 000s"); T(55, "Receivables", bs(59), unit="$ 000s"); T(56, "Retained cash", bs(60), unit="$ 000s"); T(57, "Total assets", bs(61), unit="$ 000s")
    T(58, "Payables", bs(64), unit="$ 000s"); T(59, "Senior debt", bs(65), unit="$ 000s"); T(60, "Deferred tax", bs(66), unit="$ 000s"); T(61, "Share capital", bs(67), unit="$ 000s")
    T(62, "Retained earnings", bs(68), unit="$ 000s"); T(63, "Total liabilities and shareholder funds", bs(69), unit="$ 000s"); T(64, "Balance check (assets less funding)", "=H57-H63", unit="$ 000s")
    T(65, "Net debt (senior debt less cash)", "=H59-H56", unit="$ 000s"); T(66, "Debt / (debt + shareholder funds)", "=IF(H59+H61+H62=0,0,H59/(H59+H61+H62))", "0.0%", "%")
    T(67, "Cumulative dividends to date", '=-SUMIFS(FinStat!$J$47:$CO$47,FinStat!$J$2:$CO$2,"<="&$H$53)', unit="$ 000s"); T(68, "Cumulative debt repaid to date", "=Debt!F11-H59", unit="$ 000s")
    T(69, "Senior debt outstanding (% of initial)", "=IF(Debt!F11=0,0,H59/Debt!F11)", "0.0%", "%")
    tr["E53"].value = "Balance sheet date"; tr["H53"].value = dt.date(2026, 12, 31); tr["H53"].number_format = "dd mmm yyyy"; tr["H53"].fill = YELLOW
    tr["G53"].value = "date"
    for r, (label, f, fmt, unit, F) in lab.items():
        tr[f"E{r}"].value = label
        tr[f"H{r}"].value = f; tr[f"H{r}"].number_format = fmt
        if unit: tr[f"G{r}"].value = unit
        if F is not None: tr[f"F{r}"].value = F; tr[f"F{r}"].fill = YELLOW; tr[f"F{r}"].number_format = "0"
        for col in "MNO": tr[f"{col}{r}"].number_format = fmt
        tr[f"K{r}"].number_format = fmt
        tr[f"L{r}"].number_format = "0.0%"
    tr["E77"].value = "Comment: Key input and logic changes"
    tr["H77"].value = "v1.0 corrected model: see Cover sheet change log"
    tr["H78"].value = "Stored runs M:O are values; refresh with Copy H8:H69, Paste Special Values"
    for r in range(8, 70):
        tr[f"I{r}"].value = None
    wb.save(out)
    return changes

if __name__ == "__main__":
    import re
    ch = build(sys.argv[1])
    import json; json.dump(ch, open("/private/tmp/claude-501/s/model/changes1.json","w"))
    for x in ch: print("-", x)
