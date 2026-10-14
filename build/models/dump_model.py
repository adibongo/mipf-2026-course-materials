import openpyxl, json, sys, datetime as dt, warnings, os
from openpyxl.worksheet.formula import ArrayFormula
warnings.filterwarnings("ignore")
def dump(xlsx, valsjson, out, maxrow=130):
    wf = openpyxl.load_workbook(xlsx)
    V = json.load(open(valsjson)) if valsjson and os.path.exists(valsjson) else {}
    fname = os.path.basename(xlsx)
    def v(sheet, addr):
        return V.get(f"'[{fname}]{sheet.upper()}'!{addr}")
    cols = [("A",1),("B",2),("C",3),("D",4),("E",5),("F",6),("G",7),("H",8),("J",10),("K",11),("L",12),("M",13),("N",14),("O",15),("P",16),("Q",17),("S",19),("AY",51),("CO",93)]
    out_d = {}
    for ws in wf:
        if ws.title in ("Temp", "Cover", "Dashboard"): continue
        rows = {}
        for r in range(1, min(ws.max_row, maxrow if ws.title != "InpC" else 160) + 1):
            d = {}
            for key, c in cols:
                raw = ws.cell(r, c).value
                if isinstance(raw, ArrayFormula): raw = "{" + raw.text + "}"
                if isinstance(raw, (dt.datetime, dt.date)): raw = raw.isoformat()
                if raw not in (None, ""): d[key] = raw
                val = v(ws.title, f"{key}{r}")
                if val is None and isinstance(raw, (int, float)): val = raw
                if val not in (None, ""): d[key + "_v"] = val
            if d: rows[str(r)] = d
        out_d[ws.title] = rows
    json.dump(out_d, open(out, "w"), default=str)
    return out_d
if __name__ == "__main__":
    d = dump(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None, sys.argv[3])
    print({k: len(v) for k, v in d.items()})
