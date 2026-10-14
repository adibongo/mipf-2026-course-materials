"""Python replica of TrainStat start model (Upside scenario) + project-finance extension used in Days 6-7."""
import openpyxl, datetime as dt, calendar, warnings
warnings.filterwarnings("ignore")
SRC="/Users/adibongo/Library/CloudStorage/OneDrive-Personal/Mastering Infrastructure Project Finance 2023/TrainStat_StartModel - Copy.xlsx"
def eom(d,m):
    y=d.year+(d.month-1+m)//12; mo=(d.month-1+m)%12+1
    return dt.date(y,mo,calendar.monthrange(y,mo)[1])
def addm(d,m):
    y=d.year+(d.month-1+m)//12; mo=(d.month-1+m)%12+1
    return dt.date(y,mo,1)
SRC_OVERRIDE=[None]
_CACHE={}
def load_inputs():
    key=SRC_OVERRIDE[0] or SRC
    if key not in _CACHE: _CACHE[key]=openpyxl.load_workbook(key,data_only=True)
    wb=_CACHE[key]; c=wb["InpC"]; s=wb["InpS"]
    col={"Base":10,"Downside":11,"Upside":12}
    return wb,c,s,col
def run(scn="Upside", profile=1, trains_override=None, nca0=None, fare_mult=1.0, occ_mult=1.0, om_mult=1.0, rate_add=0.0, cost_mult=1.0, fixed=False):
    wb,c,s,col=load_inputs(); k=col[scn]
    g=lambda r:c.cell(r,k).value
    N=84; first=g(9).date() if hasattr(g(9),'date') else g(9)
    acq=g(10).date(); nyrs=g(11); mpp=g(12)
    # time
    beg=[];end=[]
    for i in range(N):
        b=addm(first,i*mpp); beg.append(b); end.append(eom(b,mpp-1))
    days=[(end[i]-beg[i]).days for i in range(N)]
    acqf=[1 if end[i]==acq else 0 for i in range(N)]
    lastfc=eom(acq+dt.timedelta(days=1),nyrs*12-1)
    fcf=[1 if (end[i]>acq and end[i]<=lastfc) else 0 for i in range(N)]
    lastf=[1 if end[i]==lastfc else 0 for i in range(N)]
    lastdep=eom(acq,g(65)*12)
    depf=[1 if (end[i]>acq and end[i]<=lastdep) else 0 for i in range(N)]
    # escalation
    rates=[(g(r-0).date() if False else c.cell(r,5).value, g(r)) for r in range(17,22)]
    def infl(d):
        v=None
        for sd,rt in rates:
            if d>=sd.date(): v=rt
        return v
    per=[(1+infl(end[i]))**(1/4)-1 for i in range(N)]
    cpi=[0.0]*N
    for i in range(N):
        cpi[i]=1 if acqf[i] else (cpi[i-1]*(1+per[i]) if i>0 else 0)
    hesc=[g(23)/g(25)*x for x in cpi]; oesc=[g(23)/g(27)*x for x in cpi]
    N0=nca0 if nca0 is not None else g(64)
    # inputs
    gs=lambda r:[s.cell(r,10+i).value or 0 for i in range(N)]
    td,ti,ad,ai,od,oi=gs(8),gs(9),gs(11),gs(12),gs(14),gs(15)
    od=[x*occ_mult for x in od]; oi=[x*occ_mult for x in oi]
    prof={1:gs(18),2:gs(19),3:gs(20)}[profile]
    arr_d=[td[i]*ad[i] if fcf[i] else 0 for i in range(N)]; arr_i=[ti[i]*ai[i] if fcf[i] else 0 for i in range(N)]
    arr=[a+b for a,b in zip(arr_d,arr_i)]
    pax=[g(31)*od[i]*arr_d[i]+g(33)*oi[i]*arr_i[i] for i in range(N)]
    rev=[g(35)*fare_mult/1000*pax[i]*cpi[i] if fcf[i] else 0 for i in range(N)]
    ar=[0]*N; rec=[0]*N
    for i in range(N):
        b=ar[i-1] if i>0 else 0
        rec[i]=b+((days[i]-g(37))/days[i])*rev[i]
        ar[i]=0 if acqf[i] else b+rev[i]-rec[i]
    hv=[g(44)/1000*arr[i]*hesc[i] if fcf[i] else 0 for i in range(N)]
    hf=[g(42)/4*hesc[i] if fcf[i] else 0 for i in range(N)]
    om=[g(48)*om_mult*((1+(g(46) or 0)) if fixed else 1)/4*oesc[i] if fcf[i] else 0 for i in range(N)]
    cost=[(hv[i]+hf[i]+om[i])*cost_mult for i in range(N)]
    ap=[0]*N; paid=[0]*N
    for i in range(N):
        b=ap[i-1] if i>0 else 0
        paid[i]=b+((days[i]-g(50))/days[i])*cost[i]
        ap[i]=0 if acqf[i] else b+cost[i]-paid[i]
    ebitda=[rev[i]-cost[i] for i in range(N)]
    # NCA
    capex=[0]*N
    cx={c.cell(r,5).value.date():g(r) for r in range(56,63)}
    for i in range(N): capex[i]=cx.get(end[i],0)
    pl=g(67)*4
    cnt=list(range(1,N+1)); ret=[sum(capex[j] for j in range(i+1) if cnt[j]==cnt[i]-pl) for i in range(N)]
    cdb=[0]*N; dcap=[0]*N
    for i in range(N):
        b=cdb[i-1] if i>0 else 0
        dcap[i]=b/pl; cdb[i]=b+capex[i]-ret[i]
    d0=[N0/(g(65)*4) if depf[i] else 0 for i in range(N)]
    dep=[d0[i]+dcap[i] for i in range(N)]
    nca=[0]*N
    for i in range(N):
        b=nca[i-1] if i>0 else 0
        nca[i]=N0 if acqf[i] else b+capex[i]-dep[i]
    # tax
    twdv=[0]*N; tdep=[0]*N
    for i in range(N):
        b=twdv[i-1] if i>0 else 0
        tdep[i]=b if lastf[i] else g(84)*b
        twdv[i]=N0 if acqf[i] else b+capex[i]-tdep[i]
    # debt
    D0=N0*g(72); rate=g(77)+rate_add; dc=g(79)
    debt=[0]*N; rep=[0]*N; intr=[0]*N
    for i in range(N):
        b=debt[i-1] if i>0 else 0
        rep[i]=D0*prof[i]; intr[i]=b*rate/dc*days[i]
        debt[i]=D0 if acqf[i] else b-rep[i]
    taxable=[ebitda[i]-tdep[i]-(intr[i] if fixed else 0) for i in range(N)]
    taxpaid=[g(86)*x for x in taxable]
    pbt=[ebitda[i]-dep[i]-intr[i] for i in range(N)]
    taxexp=[g(86)*x for x in pbt]; pat=[pbt[i]-taxexp[i] for i in range(N)]
    cfads=[rec[i]-paid[i]-capex[i]-taxpaid[i] for i in range(N)]
    sc0=N0*(1-g(72)); sc=[0]*N; red=[0]*N
    cfd=[0]*N; div=[0]*N; rcash=[0]*N; re=[0]*N
    for i in range(N):
        scb=sc[i-1] if i>0 else 0
        red[i]=scb*lastf[i]; sc[i]=sc0 if acqf[i] else scb-red[i]
        cfs=cfads[i]-intr[i]-rep[i]; cfd[i]=cfs-red[i]
        rcb=rcash[i-1] if i>0 else 0; reb=re[i-1] if i>0 else 0
        avail_c=max(rcb+cfd[i],0); avail_e=max(reb+pat[i],0)
        div[i]=min(avail_c,avail_e)
        rcash[i]=0 if acqf[i] else rcb+cfd[i]-div[i]
        re[i]=0 if acqf[i] else reb+pat[i]-div[i]
    r=g(93); vd=g(90).date() if hasattr(g(90),'date') else g(90)
    yrs=[(end[i]-vd).days/g(91) for i in range(N)]; df=[1/(1+r)**y for y in yrs]
    sh=[(div[i]+red[i])-(sc0 if acqf[i] else 0) for i in range(N)]
    npv=sum(sh[i]*df[i] for i in range(N))
    return dict(N=N,beg=beg,end=end,days=days,acq=acqf,fc=fcf,lastf=lastf,cpi=cpi,rev=rev,cost=cost,ebitda=ebitda,dep=dep,tdep=tdep,
        capex=capex,intr=intr,rep=rep,debt=debt,pbt=pbt,pat=pat,taxpaid=taxpaid,taxexp=taxexp,cfads=cfads,div=div,sh=sh,df=df,npv=npv,
        arr=arr,pax=pax,hv=hv,hf=hf,om=om,rec=rec,ar=ar,ap=ap,paid=paid,nca=nca,twdv=twdv,sc=sc,red=red,rcash=rcash,re=re,D0=D0,rate=rate,dc=dc,g=g,prof=prof,arr_d=arr_d,arr_i=arr_i,hesc=hesc,oesc=oesc)
def xirr(vals,dates,lo=-0.9,hi=10):
    def f(r): return sum(v/(1+r)**((d-dates[0]).days/365) for v,d in zip(vals,dates))
    for _ in range(200):
        m=(lo+hi)/2
        if f(lo)*f(m)<=0: hi=m
        else: lo=m
    return m
if __name__=="__main__":
    m=run()
    S=lambda k:sum(m[k])
    print("rev",S("rev"),"cost",S("cost"),"ebitda",S("ebitda"),"dep",S("dep"),"int",S("intr"),"pbt",S("pbt"),"tax exp",S("taxexp"),"pat",S("pat"),"cfads",S("cfads"))
    i0=[i for i in range(m["N"]) if m["acq"][i]][0]
    vals=[m["sh"][i] if m["sh"][i]!=0 else (0 if i!=0 else -0.001) for i in range(m["N"])]
    print("npv",m["npv"],"irr",xirr(vals,m["end"]))
