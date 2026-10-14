# -*- coding: utf-8 -*-
"""Stage 2 and 3: PF extension (Constr, Sculpt, DSRA, SHL, Ratios) + Cover + Dashboard. Usage: build_final.py start|pf out.xlsx"""
import re, sys, copy, json, datetime as dt, warnings
warnings.filterwarnings("ignore")
sys.path.insert(0, "/private/tmp/claude-501/s/model")
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter as L
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.worksheet.hyperlink import Hyperlink
from fixlib import *

NAVY, GOLD, BLUE, TEAL = "141B4D", "D99A2B", "0B5E82", "1C7293"
YELLOW = PatternFill("solid", fgColor="FFFF99")
LOGO = "/Users/adibongo/Library/CloudStorage/OneDrive-Personal/Mastering Infrastructure Project Finance 2023/MIPF-October-2026/website/assets/logo/"
FMT = {"%": "0.00%", "% p.a": "0.00%", "x": '0.00"x"', "date": "dd mmm yy", "flag": "0", "counter": "0", "check": "#,##0.000_);(#,##0.000);\"-  \"", "days": "0", "periods": "0", "quarters": "0", "text": "General", "switch": "0"}
NUM = '#,##0_);(#,##0);"-  "'
FIN = [(15, 1), (43, 3), (61, 1), (66, 1)]  # FinStat insertions (stage-1 numbering), applied in descending order

def FS(r):  # final FinStat row for a stage-1 row
    return r + sum(n for at, n in FIN if r >= at)
def DB(r): return r + (1 if r >= 17 else 0)
def VL(r): return r + (2 if r >= 23 else 0)

def add_ts(wb, name, tab, ref="OpRev", main_ruler=True):
    ws = wb.create_sheet(name)
    src = wb[ref]
    for k, d in src.column_dimensions.items():
        nd = ws.column_dimensions[k]; nd.width = d.width; nd.min = d.min; nd.max = d.max; nd.hidden = d.hidden
    ws.sheet_view.showGridLines = False; ws.sheet_view.zoomScale = src.sheet_view.zoomScale or 100
    ws.freeze_panes = "J6"
    ws.sheet_properties.tabColor = tab
    for r in range(1, 6):
        for c in range(1, 94):
            s = src.cell(r, c)
            if s.has_style: ws.cell(r, c)._style = copy.copy(s._style)
            if main_ruler and s.value is not None: ws.cell(r, c).value = s.value
    return ws

def style_like(wb, ws, r, kind):
    ref = {"sec": 7, "sub": 9, "pull": 37, "tlpull": 11, "calc": 19, "res": 42, "total": 21}[kind]
    copy_row_style(wb["OpRev"], ref, r) if False else None
    src = wb["OpRev"]
    for c in range(1, 94):
        s = src.cell(ref, c)
        if s.has_style: ws.cell(r, c)._style = copy.copy(s._style)

def put(wb, ws, spec, ncols=84, firstcol=10):
    for (r, E, F, G, Hf, Jf, key) in spec:
        pure_pull = isinstance(E, str) and E.startswith("=") and Jf is None
        kind = "pull" if pure_pull else ("tlpull" if (Jf and re.fullmatch(r"=\w+!\w\$?\d+", Jf)) else ("res" if key in RES else "calc"))
        style_like(wb, ws, r, kind)
        if E is not None: ws.cell(r, 5).value = E
        if F is not None: ws.cell(r, 6).value = F
        if G is not None: ws.cell(r, 7).value = G
        if Hf: ws.cell(r, 8).value = Hf
        if Jf: fill_row(ws, r, Jf, firstcol, firstcol + ncols - 1)
        unit = G if isinstance(G, str) else ""
        fmt = FMT.get(unit, NUM)
        if unit.startswith("="): fmt = NUM
        for c in range(6, firstcol + ncols):
            if c in (7,): continue
            ws.cell(r, c).number_format = fmt
        if unit == "date" and Jf: pass
RES = {"cap", "size", "uses", "debt", "eq", "end", "shl0", "sc0", "nca0", "d0", "e0", "bend", "ds", "llcr0", "plcr0", "mind", "avgd", "dscr", "llcr", "plcr", "tgt"}

def sec(wb, ws, r, text, kind="sec", col="A"):
    style_like(wb, ws, r, kind); ws[f"{col}{r}"].value = text

def inpc_row(wb, r, label, vals, unit, fmt, header=None):
    c = wb["InpC"]
    src = {"%": 77, "x": 77, "num": 42, "int": 31, "date": 9}[fmt if fmt in ("%", "x", "int", "date") else "num"]
    copy_row_style(c, src, r, 93)
    c[f"E{r}"].value = label; c[f"G{r}"].value = unit
    for col, v in zip("JKL", vals): c[f"{col}{r}"].value = v
    c[f"F{r}"].value = ArrayFormula(f"F{r}", f"=INDEX(I{r}:M{r},$F$4)")
    nf = {"%": "0.00%", "x": '0.00"x"', "int": "#,##0", "num": NUM}[fmt]
    for col in "FJKL": c[f"{col}{r}"].number_format = nf
    for col in "JKL": c[f"{col}{r}"].fill = YELLOW

def build_pf(wb, sol):
    c = wb["InpC"]
    # ------------------------------------------------------------ insertions (descending)
    insert_rows(wb, "FinStat", 66, 1, style_from=65)
    insert_rows(wb, "FinStat", 61, 1, style_from=60)
    insert_rows(wb, "FinStat", 43, 3, style_from=41)
    insert_rows(wb, "FinStat", 15, 1, style_from=14)
    insert_rows(wb, "Debt", 17, 1, style_from=16)
    insert_rows(wb, "Val", 23, 2, style_from=22)
    fs, db, val, eq, tax = wb["FinStat"], wb["Debt"], wb["Val"], wb["Equity"], wb["Tax"]
    # ------------------------------------------------------------ InpC inputs
    sec(wb, c, 111, "CONSTRUCTION AND FUNDING", "sec");
    for col in range(1, 8): c.cell(111, col)._style = copy.copy(c.cell(29, col)._style)
    c["A111"].value = "CONSTRUCTION AND FUNDING"
    gear = {k: sol[k]["g"] for k in ("Base", "Downside", "Upside")}
    inpc_row(wb, 113, "Construction period", [8, 8, 8], "quarters", "int")
    inpc_row(wb, 115, "Total construction cost (hard cost)", [94850] * 3, "$ 000s", "num")
    for i, p in enumerate([0.05, 0.10, 0.15, 0.20, 0.20, 0.15, 0.10, 0.05]):
        inpc_row(wb, 118 + i, i + 1, [p] * 3, "% of cost", "%"); c[f"E{118+i}"].number_format = "0"
    c["E117"].value = "Construction spend profile (quarter number and share of cost)"
    inpc_row(wb, 127, "Upfront fee (% of facility)", [0.02] * 3, "%", "%")
    inpc_row(wb, 129, "Commitment fee (% p.a. on undrawn facility)", [0.0075] * 3, "% p.a", "%")
    inpc_row(wb, 131, "Senior debt facility (commitment)", [sol[k]["fac"] for k in ("Base", "Downside", "Upside")], "$ 000s", "num")
    inpc_row(wb, 133, "Drawdown method (1 pro-rata, 2 equity first)", [1] * 3, "switch", "int")
    inpc_row(wb, 135, "Equity commitment (used by method 2)", [40000] * 3, "$ 000s", "num")
    for col in range(1, 8): c.cell(137, col)._style = copy.copy(c.cell(29, col)._style)
    c["A137"].value = "SENIOR DEBT SIZING"
    inpc_row(wb, 139, "Senior debt tenor", [40] * 3, "quarters", "int")
    inpc_row(wb, 140, "Target DSCR for sculpting", [2.25] * 3, "x", "x")
    inpc_row(wb, 141, "Maximum gearing (debt / total uses)", [0.75] * 3, "%", "%")
    inpc_row(wb, 143, "DSRA cover (periods of debt service)", [2] * 3, "periods", "int")
    for col in range(1, 8): c.cell(145, col)._style = copy.copy(c.cell(29, col)._style)
    c["A145"].value = "SHAREHOLDER LOAN"
    inpc_row(wb, 147, "Shareholder loan share of sponsor funding", [0.60] * 3, "%", "%")
    inpc_row(wb, 148, "Shareholder loan interest rate", [0.10] * 3, "% p.a", "%")
    for col, k in zip("JKL", ("Base", "Downside", "Upside")): c[f"{col}72"].value = gear[k]     # gearing solved by Goal Seek, per scenario
    for col in "JKL": c[f"{col}74"].value = 4        # sculpted profile
    c["N72"].value = "Solved per scenario with Goal Seek: Sculpt!F32 (debt drawn less debt sized) = 0 by changing the active scenario's cell"
    # NCA initial balance now comes from construction
    for col in "JKL": c[f"{col}64"].value = "=Constr!$F$52"
    c["N64"].value = "Linked to Constr!F52: capitalised project cost (hard cost + fees + interest during construction)"
    # ------------------------------------------------------------ Constr sheet (own 8-quarter timeline)
    ws = add_ts(wb, "Constr", TEAL, main_ruler=False)
    ws["A1"].value = '= RIGHT(CELL("filename", A1), LEN(CELL("filename", A1)) - SEARCH("]", CELL("filename", A1)))'
    ws["E2"].value = "Construction period ending"; ws["E3"].value = "Phase"; ws["E4"].value = "Financial year"; ws["E5"].value = "Construction quarter counter"
    ws["F5"].value = "Constant"; ws["G5"].value = "Unit"; ws["H5"].value = "Total"
    for k in range(8):
        col = 10 + k; cl = L(col)
        ws.cell(2, col).value = f"={cl}13"; ws.cell(3, col).value = "Construction"; ws.cell(4, col).value = f"=YEAR({cl}13)"; ws.cell(5, col).value = f"={L(col-1)}5+1" if k else 1
        ws.cell(2, col).number_format = "dd mmm yy"
    cons = [
     (9, "=InpC!E$10", "=InpC!F$10", "=InpC!G$10", None, None, "cod"), (10, "=InpC!E$12", "=InpC!F$12", "=InpC!G$12", None, None, "mpp"), (11, "=InpC!E$113", "=InpC!F$113", "=InpC!G$113", None, None, "nq"),
     (13, "Construction period ending", None, "date", None, "=EOMONTH($F$9,-$F$10*($F$11-J$5))", "end"), (14, "Days in construction period", None, "days", "=SUM(J14:Q14)", "=J13-EOMONTH(J13,-$F$10)", "days"),
     (17, "=InpC!E$115", "=InpC!F$115", "=InpC!G$115", None, None, "cost"), (18, "Spend profile", None, "%", "=SUM(J18:Q18)", "=INDEX(InpC!$F$118:$F$125,J$5)", "prof"), (19, "Construction cost spend", None, "$ 000s", "=SUM(J19:Q19)", "=$F$17*J18", "spend"),
     (22, "=InpC!E$127", "=InpC!F$127", "=InpC!G$127", None, None, "upf"), (23, "=InpC!E$131", "=InpC!F$131", "=InpC!G$131", None, None, "fac"), (24, "=InpC!E$129", "=InpC!F$129", "=InpC!G$129", None, None, "cfp"),
     (25, "Upfront fee", None, "$ 000s", "=SUM(J25:Q25)", "=IF(J$5=1,$F$22*$F$23,0)", "upfee"), (28, "=InpC!E$77", "=InpC!F$77", "=InpC!G$77", None, None, "rate"), (29, "=InpC!E$79", "=InpC!F$79", "=InpC!G$79", None, None, "dc"),
     (30, "Senior debt balance BEG", None, "$ 000s", None, "=I49", "beg"), (31, "Interest during construction (IDC)", None, "$ 000s", "=SUM(J31:Q31)", "=J30*$F$28/$F$29*J14", "idc"),
     (32, "Commitment fee", None, "$ 000s", "=SUM(J32:Q32)", "=($F$23-J30)*$F$24/$F$29*J14", "cfee"),
     (35, "=E$19", None, "$ 000s", "=SUM(J35:Q35)", "=J19", "u1"), (36, "=E$25", None, "$ 000s", "=SUM(J36:Q36)", "=J25", "u2"), (37, "=E$31", None, "$ 000s", "=SUM(J37:Q37)", "=J31", "u3"), (38, "=E$32", None, "$ 000s", "=SUM(J38:Q38)", "=J32", "u4"),
     (39, "Total uses of funds", None, "$ 000s", "=SUM(J39:Q39)", "=SUM(J35:J38)", "uses"),
     (42, "=InpC!E$72", "=InpC!F$72", "=InpC!G$72", None, None, "gear"), (43, "=InpC!E$133", "=InpC!F$133", "=InpC!G$133", None, None, "meth"), (44, "=InpC!E$135", "=InpC!F$135", "=InpC!G$135", None, None, "eqc"),
     (45, "Equity drawdown", None, "$ 000s", "=SUM(J45:Q45)", "=IF($F$43=1,J39*(1-$F$42),MIN(J39,$F$44-SUM($I45:I45)))", "eq"), (46, "Senior debt drawdown", None, "$ 000s", "=SUM(J46:Q46)", "=J39-J45", "debt"),
     (47, "Total sources of funds", None, "$ 000s", "=SUM(J47:Q47)", "=J45+J46", "src"), (48, "Sources less uses (check)", None, "check", "=SUM(J48:Q48)", "=J47-J39", "chk"), (49, "Senior debt balance END", None, "$ 000s", None, "=J30+J46", "end_bal"),
     (52, "Capitalised project cost at COD", "=H39", "$ 000s", None, None, "nca0"), (53, "Senior debt at COD", "=Q49", "$ 000s", None, None, "d0"), (54, "Sponsor funding at COD (equity)", "=H45", "$ 000s", None, None, "e0"), (55, "Check: cost less debt less equity", "=F52-F53-F54", "check", None, None, "c0")]
    for r, t in ((7, "CONSTRUCTION TIMELINE"), (16, "CONSTRUCTION COST"), (21, "FEES"), (27, "INTEREST AND COMMITMENT FEE DURING CONSTRUCTION"), (34, "USES OF FUNDS"), (41, "SOURCES OF FUNDS"), (51, "OPENING BALANCES AT COD (financial close)")):
        sec(wb, ws, r, t)
    put(wb, ws, cons, 8)
    for r in (13,):
        for col in range(10, 18): ws.cell(r, col).number_format = "dd mmm yy"
    for r in (9,):
        ws.cell(r, 6).number_format = "dd mmm yy"
    # ------------------------------------------------------------ Debt: profile 4 and opening debt
    db["E17"].value = "Senior debt profile 4 - sculpted"; db["G17"].value = "%"; db["H17"].value = "= SUM(J17:CO17)"
    fill_row(db, 17, "=Sculpt!J$40")
    for col in range(10, 94): db.cell(17, col).number_format = "0.00%"
    db["E18"].value = ArrayFormula("E18", '=INDEX(E14:E17,F13)&" - active"')
    fill_row(db, 18, "=INDEX(J14:J17,$F13)")
    db["F11"].value = "=Constr!$F$53"
    db["F10"].value = "=InpC!F$72"
    eq["F16"].value = "=SHL!$F$14"
    # ------------------------------------------------------------ tax: interest incl. shareholder loan, and unlevered tax for sizing
    tax["E26"].value = "Interest due (senior debt and shareholder loan)"; tax["F26"].value = None; tax["G26"].value = "=SMU"
    fill_row(tax, 26, f"=Debt!J${DB(32)}+SHL!J$20")
    for r, lab, f in ((52, "Taxable income before interest deduction", "=J24-J25"), (53, "Tax paid before interest deduction POS (for debt sizing)", "=$F29*J52")):
        copy_row_style(tax, 27, r); tax[f"E{r}"].value = lab; tax[f"G{r}"].value = "=SMU"; tax[f"H{r}"].value = f"= SUM(J{r}:CO{r})"
        fill_row(tax, r, f)
    sec(wb, tax, 51, "TAX BEFORE INTEREST DEDUCTION (used only for sizing debt)", "sec")
    # ------------------------------------------------------------ Sculpt
    ws = add_ts(wb, "Sculpt", GOLD)
    cf_row = FS(35)
    sculpt = [
     (9, "=InpC!E$139", "=InpC!F$139", "=InpC!G$139", None, None, "ten"), (10, "Forecast period flag", None, "flag", None, "=Time!J$69", "fc"), (11, "Operating period counter", None, "counter", None, "=I11+J10", "cnt"),
     (12, "Repayment period flag", None, "flag", "=SUM(J12:CO12)", "=IF(AND(J11>=1,J11<=$F$9),1,0)", "rflag"),
     (15, "=InpC!E$140", "=InpC!F$140", "=InpC!G$140", None, None, "dscr_t"), (16, "=InpC!E$77", "=InpC!F$77", "=InpC!G$77", None, None, "rate"), (17, "=InpC!E$79", "=InpC!F$79", "=InpC!G$79", None, None, "dc"),
     (18, "Days in period", None, "days", None, "=Time!J$30", "days"), (19, "Periodic interest rate", None, "%", None, "=$F$16/$F$17*J18", "pr"),
     (20, "CFADS for sizing (tax before interest deduction)", None, "$ 000s", "=SUM(J20:CO20)", f"=FinStat!J${cf_row}-Tax!J$53", "cfads"), (21, "CFADS available to lenders", None, "$ 000s", "=SUM(J21:CO21)", "=J20*J12", "cf_l"),
     (22, "Debt service at target DSCR", None, "$ 000s", "=SUM(J22:CO22)", "=J21/$F$15", "ds_t"), (23, "Debt capacity balance END", None, "$ 000s", None, "=(K23+K22)/(1+K19)", "capbal"),
     (25, "Acquisition / COD flag", None, "flag", None, "=Time!J$45", "acq"), (26, "Debt capacity at COD", "=SUMPRODUCT(J23:CO23,J25:CO25)", "$ 000s", None, None, "cap"),
     (27, "=InpC!E$141", "=InpC!F$141", "=InpC!G$141", None, None, "gmax"), (28, "Debt limit from gearing", "=F27*Constr!F52", "$ 000s", None, None, "glim"), (29, "Senior debt size", "=MIN(F26,F28)", "$ 000s", None, None, "size"),
     (30, "Binding constraint", '=IF(F26<=F28,"DSCR","Gearing")', "text", None, None, "bind"), (31, "Sculpted DSCR at debt size", "=F15*F26/F29", "x", None, None, "sdscr"), (32, "Debt drawn less debt sized (Goal Seek target = 0)", "=Constr!F53-F29", "check", None, None, "gs"),
     (34, "Sculpted debt service", None, "$ 000s", "=SUM(J34:CO34)", "=J21/$F$31", "ds"), (35, "Senior debt balance BEG", None, "$ 000s", None, "=I38", "bbeg"), (36, "Interest", None, "$ 000s", "=SUM(J36:CO36)", "=J35*J19", "int"),
     (37, "Principal", None, "$ 000s", "=SUM(J37:CO37)", "=J34-J36", "prin"), (38, "Senior debt balance END", None, "$ 000s", None, "=IF(J25=1,$F$29,J35-J37)", "bend"), (40, "Sculpted repayment profile (% of opening debt)", None, "%", "=SUM(J40:CO40)", "=J37/$F$29", "profile")]
    for r, t in ((7, "LOAN TIMELINE"), (14, "DEBT SIZING"), (33, "SCULPTED REPAYMENT")): sec(wb, ws, r, t)
    put(wb, ws, sculpt)
    for r in (19,):
        for col in range(10, 94): ws.cell(r, col).number_format = "0.0000%"
    ws["F30"].number_format = "General"
    # ------------------------------------------------------------ DSRA
    ws = add_ts(wb, "DSRA", BLUE)
    cfa = FS(42)
    dsra = [
     (9, "=InpC!E$143", "=InpC!F$143", "=InpC!G$143", None, None, "n"), (10, "Operating period counter", None, "counter", None, "=Sculpt!J$11", "cnt"),
     (11, "Senior debt service", None, "$ 000s", "=SUM(J11:CO11)", f"=FinStat!J${FS(88)}", "ds"),
     (12, "DSRA target (next periods of debt service)", None, "$ 000s", None, '=SUMIFS($J11:$CO11,$J10:$CO10,">"&J10,$J10:$CO10,"<="&J10+$F$9)*(1-Time!J$58+Time!J$45)', "tgt"),
     (13, "Cash flow after senior debt service", None, "$ 000s", None, f"=FinStat!J${cfa}", "cf"), (14, "DSRA balance BEG", None, "$ 000s", None, "=I18", "beg"),
     (15, "Funding", None, "$ 000s", "=SUM(J15:CO15)", "=MIN(MAX(J12-J14,0),MAX(J13,0))", "fund"), (16, "Release", None, "$ 000s", "=SUM(J16:CO16)", "=MAX(J14-J12,0)", "rel"),
     (17, "DSRA (funding) / release", None, "$ 000s", "=SUM(J17:CO17)", "=J16-J15", "net"), (18, "DSRA balance END", None, "$ 000s", None, "=IF(Time!J$45=1,0,J14+J15-J16)", "end")]
    sec(wb, ws, 7, "DEBT SERVICE RESERVE ACCOUNT"); put(wb, ws, dsra)
    # ------------------------------------------------------------ SHL
    ws = add_ts(wb, "SHL", TEAL)
    shl = [
     (9, "=InpC!E$147", "=InpC!F$147", "=InpC!G$147", None, None, "pct"), (10, "=InpC!E$148", "=InpC!F$148", "=InpC!G$148", None, None, "rate"), (11, "=InpC!E$79", "=InpC!F$79", "=InpC!G$79", None, None, "dc"),
     (12, "Sponsor funding at COD", "=Constr!F54", "$ 000s", None, None, "sp"), (13, "Shareholder loan - initial balance", "=F12*F9", "$ 000s", None, None, "shl0"), (14, "Share capital - initial balance", "=F12-F13", "$ 000s", None, None, "sc0"),
     (16, "Cash available for shareholder loan", None, "$ 000s", None, f"=MAX(FinStat!J${cfa}+DSRA!J$17,0)", "av"), (17, "Post-senior-loan flag", None, "flag", None, "=IF(Sculpt!J$11>Sculpt!$F$9,1,0)", "post"),
     (18, "Days in period", None, "days", None, "=Time!J$30", "days"), (19, "Shareholder loan balance BEG", None, "$ 000s", None, "=I24", "beg"),
     (20, "Interest due", None, "$ 000s", "=SUM(J20:CO20)", "=J19*$F$10/$F$11*J18", "due"), (21, "Interest paid", None, "$ 000s", "=SUM(J21:CO21)", "=MIN(J20,J16)", "paid"),
     (22, "Interest capitalised (unpaid)", None, "$ 000s", "=SUM(J22:CO22)", "=J20-J21", "capz"), (23, "Principal repaid", None, "$ 000s", "=SUM(J23:CO23)", "=J17*MIN(J19+J22,J16-J21)", "prin"),
     (24, "Shareholder loan balance END", None, "$ 000s", None, "=IF(Time!J$45=1,$F$13,J19+J22-J23)", "end"),
     (25, "Shareholder loan interest", None, "$ 000s", "=SUM(J25:CO25)", "=-1*J20", "is"), (26, "Shareholder loan interest paid", None, "$ 000s", "=SUM(J26:CO26)", "=-1*J21", "cf1"), (27, "Shareholder loan principal repaid", None, "$ 000s", "=SUM(J27:CO27)", "=-1*J23", "cf2")]
    sec(wb, ws, 7, "SHAREHOLDER LOAN"); put(wb, ws, shl)
    ws["E25"].font = copy.copy(ws["E24"].font)
    # ------------------------------------------------------------ FinStat integration
    def fsrow(r, label_src, jf, H=True, style=None):
        copy_row_style(fs, style or 14, r)
        fs[f"E{r}"].value = label_src; fs[f"G{r}"].value = "$ 000s"
        if H: fs[f"H{r}"].value = f"= SUM(J{r}:CO{r})"
        fill_row(fs, r, jf)
    fsrow(15, "=SHL!E$25", "=SHL!J$25")
    fs["E15"].value = "=SHL!E$25"
    fill_row(fs, 16, "=J11+SUM(J13:J15)")
    fsrow(44, "=DSRA!E$17", "=DSRA!J$17", style=41); fsrow(45, "=SHL!E$26", "=SHL!J$26", style=41); fsrow(46, "=SHL!E$27", "=SHL!J$27", style=41)
    fill_row(fs, 49, "=J43+SUM(J44:J48)")
    fsrow(65, "=DSRA!E$18", "=DSRA!J$18", H=False, style=64); fsrow(71, "=SHL!E$24", "=SHL!J$24", H=False, style=70)
    fill_row(fs, 66, "=SUM(J62:J65)"); fill_row(fs, 75, "=SUM(J69:J74)")
    # ------------------------------------------------------------ Valuation: blended sponsor view
    for r, src in ((23, 21), (24, 23)):
        copy_row_style(val, 22, r)
        val[f"E{r}"].value = f"=SHL!E${src}"; val[f"F{r}"].value = f"=SHL!F${src}"; val[f"G{r}"].value = f"=SHL!G${src}"; val[f"H{r}"].value = f"= SUM(J{r}:CO{r})"
        fill_row(val, r, f"=SHL!J${src}")
    val["E23"].value = "Shareholder loan interest received"; val["E24"].value = "Shareholder loan principal received"
    fill_row(val, 25, "=SUM(J21:J24)")
    val["E27"].value = "=SHL!E$12"; val["F27"].value = "=SHL!F$12"; val["G27"].value = "=SHL!G$12"
    # ------------------------------------------------------------ Ratios sheet
    ws = add_ts(wb, "Ratios", NAVY)
    rat = [
     (9, "Senior debt service", None, "$ 000s", "=SUM(J9:CO9)", f"=FinStat!J${FS(88)}", "ds"), (10, "CFADS", None, "$ 000s", "=SUM(J10:CO10)", f"=FinStat!J${FS(38)}", "cfads"), (11, "DSCR", None, "x", None, f"=FinStat!J${FS(90)}", "dscr"),
     (13, "Minimum DSCR", f"=FinStat!F{FS(92)}", "x", None, None, "mind"), (14, "Average DSCR", f"=FinStat!F{FS(93)}", "x", None, None, "avgd"),
     (17, "Periodic discount rate (debt cost)", None, "%", None, "=Sculpt!J$19", "pr"), (18, "CFADS in loan life", None, "$ 000s", None, f"=FinStat!J${FS(38)}*Sculpt!J$12", "cfl"),
     (19, "PV of CFADS in loan life (END of period)", None, "$ 000s", None, "=(K19+K18)/(1+K17)", "pvl"), (20, "Senior debt balance END", None, "$ 000s", None, f"=Debt!J${DB(44)}", "bal"),
     (21, "LLCR", None, "x", None, "=IF(J20>CHK_TOL,J19/J20,0)", "llcr"), (22, "LLCR at COD", "=SUMPRODUCT(J21:CO21,Time!J45:CO45)", "x", None, None, "llcr0"),
     (25, "CFADS in project life", None, "$ 000s", None, f"=FinStat!J${FS(38)}*Time!J$69", "cfp"), (26, "PV of CFADS in project life (END of period)", None, "$ 000s", None, "=(K26+K25)/(1+K17)", "pvp"),
     (27, "PLCR", None, "x", None, "=IF(J20>CHK_TOL,J26/J20,0)", "plcr"), (28, "PLCR at COD", "=SUMPRODUCT(J27:CO27,Time!J45:CO45)", "x", None, None, "plcr0")]
    for r, t in ((7, "DEBT SERVICE COVER"), (16, "LOAN LIFE COVER RATIO (LLCR)"), (24, "PROJECT LIFE COVER RATIO (PLCR)")): sec(wb, ws, r, t)
    put(wb, ws, rat)
    # order sheets
    order = ["InpC", "InpS", "Time", "Esc", "Constr", "OpRev", "OpCost", "Tax", "NCA", "Debt", "Sculpt", "DSRA", "SHL", "Equity", "FinStat", "Ratios", "Val", "Track", "Temp"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in order]
    return wb

def track_pf(wb):
    tr = wb["Track"]
    # extend the key outputs and balance sheet lines for the PF model
    tr["E8"].value = "Sponsor NPV at required return (blended)"; tr["E9"].value = "Sponsor IRR (XIRR, blended)"
    tr["H11"].value = "=Ratios!F13"; tr["H12"].value = "=Debt!F11"
    tr["E12"].value = "Senior debt at financial close"
    # replace three BS-date lines with PF items where useful
    bs = lambda row: f"=INDEX(FinStat!$J${row}:$CO${row},MATCH($H$53,FinStat!$J$2:$CO$2,0))"
    tr["E69"].value = "DSRA balance"; tr["H69"].value = bs(FS(61) + 0 if False else 65)
    tr["E68"].value = "Shareholder loan balance"; tr["H68"].value = bs(71)
    tr["E67"].value = "Cumulative dividends to date"
    tr["H66"].value = "=IF(H59+H61+H62+H68=0,0,H59/(H59+H61+H62+H68))"; tr["E66"].value = "Senior debt / (debt + shareholder funds)"
    tr["F52"].value = None
    tr["E52"].value = "Minimum LLCR at COD"; tr["H52"].value = "=Ratios!F22"; tr["G52"].value = "x"
    tr["E51"].value = "Debt capacity (DSCR test)"; tr["H51"].value = "=Sculpt!F26"; tr["G51"].value = "$ 000s"
    for r in (68, 69):
        for col in "HMNO": tr[f"{col}{r}"].number_format = NUM
    for col in "HMNO": tr[f"{col}52"].number_format = '0.00"x"'

def add_cover_dashboard(wb, kind, changes, scen_note):
    cv = wb.create_sheet("Cover", 0); ds = wb.create_sheet("Dashboard", 1)
    for ws in (cv, ds):
        ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = GOLD if ws is cv else NAVY
    navy = PatternFill("solid", fgColor=NAVY); gold = PatternFill("solid", fgColor=GOLD); mist = PatternFill("solid", fgColor="F3F6F9")
    W = Font(name="Calibri", color="FFFFFF", bold=True)
    title = "TrainStat Project Finance Model: complete" if kind == "pf" else "TrainStat Start Model: corrected and completed"
    cv.column_dimensions["A"].width = 3
    for col, w in zip("BCDEFGH", (24, 44, 22, 22, 22, 22, 12)): cv.column_dimensions[col].width = w
    for r in range(1, 6):
        for c in range(1, 9): cv.cell(r, c).fill = navy
    cv["B2"].value = title; cv["B2"].font = Font(name="Georgia", size=22, bold=True, color="FFFFFF")
    cv["B3"].value = "Mastering Infrastructure and Project Finance Boot Camp  |  Lagos Business School  |  14 to 21 October 2026"; cv["B3"].font = Font(name="Calibri", size=12, color="F0C06A")
    cv["B4"].value = "bongo.ai  |  MIPF  |  Prof Bongo Adi  |  +234 706 164 8699  |  badi@lbs.edu.ng  |  bongo.adi@fulbrightmail.org"; cv["B4"].font = Font(name="Calibri", size=11, color="C9D3F5")
    for f, anchor, h in (("lbs-logo.png", "H1", 52), ("mipf_hex_256.png", "F1", 58), ("bongo_ai_hex_256.png", "G1", 58)):
        im = XLImage(LOGO + f); ratio = im.width / im.height; im.height = h; im.width = int(h * ratio); cv.add_image(im, anchor)
    cv.row_dimensions[2].height = 30
    r = 7
    cv[f"B{r}"].value = "MODEL STATUS"; cv[f"B{r}"].font = Font(bold=True, color=BLUE)
    ck = "FinStat!F{a}+FinStat!F{b}+FinStat!F{c}+Time!F120"
    if kind == "pf": a, b, c_ = FS(74), FS(77), FS(80)
    else: a, b, c_ = 74, 77, 80
    cv[f"C{r}"].value = f'=IF({ck.format(a=a,b=b,c=c_)}+IF(ABS(Constr!F55)>CHK_TOL,1,0)+IF(SUMPRODUCT(ABS(Constr!J48:Q48))>CHK_TOL,1,0)=0,"ALL CHECKS PASSED","CHECK ERRORS: see FinStat and Constr")' if kind == "pf" else f'=IF({ck.format(a=a,b=b,c=c_)}=0,"ALL CHECKS PASSED","CHECK ERRORS: see FinStat")'
    cv[f"C{r}"].font = Font(bold=True, size=13, color="1E7B34")
    cv[f"B{r+1}"].value = "Active scenario"; cv[f"C{r+1}"].value = "=InpC!F3"; cv[f"C{r+1}"].font = Font(bold=True)
    cv[f"B{r+2}"].value = "Scenario definitions"; cv[f"C{r+2}"].value = scen_note
    cv[f"B{r+2}"].alignment = Alignment(vertical="top"); cv[f"C{r+2}"].alignment = Alignment(wrap_text=True, vertical="top"); cv.merge_cells(f"C{r+2}:G{r+2}"); cv.row_dimensions[r + 2].height = 48
    if kind == "pf":
        cv["B10"].value = "Debt sizing"; cv["C10"].value = '=IF(ABS(Sculpt!F32)<1,"Debt drawn equals debt sized","RE-RUN GOAL SEEK: debt drawn less debt sized = "&TEXT(Sculpt!F32,"#,##0"))'; cv["C10"].font = Font(bold=True, color=BLUE)
    r = 11
    cv[f"B{r}"].value = "CONTENTS"; cv[f"B{r}"].font = Font(bold=True, color=BLUE)
    sheets = [("Dashboard", "Key results, annual summary and charts"), ("InpC", "Constants and scenario switch (yellow = typed assumption)"), ("InpS", "Time-series inputs: trains, arrivals, occupancy, repayment profiles, extra funding"), ("Time", "Timeline, flags and integrity checks"),
              ("Esc", "Inflation and escalation factors")]
    if kind == "pf": sheets += [("Constr", "Construction period: spend, fees, interest during construction, sources and uses, opening balances")]
    sheets += [("OpRev", "Demand and operating revenue, cash received, receivables"), ("OpCost", "Operating cost, cash paid, payables"), ("Tax", "Tax depreciation, tax paid, tax expense, deferred tax")]
    sheets += [("NCA", "Capital expenditure, depreciation, asset balance"), ("Debt", "Senior debt: repayment profile, interest, balance")]
    if kind == "pf": sheets += [("Sculpt", "Debt capacity, sizing and sculpted repayment"), ("DSRA", "Debt service reserve account"), ("SHL", "Shareholder loan")]
    sheets += [("Equity", "Share capital, dividends, retained cash and earnings"), ("FinStat", "Income statement, cash flow cascade, balance sheet, checks, DSCR")]
    if kind == "pf": sheets += [("Ratios", "DSCR, LLCR and PLCR")]
    sheets += [("Val", "Shareholder valuation: NPV and XIRR"), ("Track", "Scenario log: live results and stored runs")]
    for i, (n, d) in enumerate(sheets):
        rr = r + 1 + i
        cv[f"B{rr}"].value = n; cv[f"B{rr}"].hyperlink = f"#'{n}'!A1"; cv[f"B{rr}"].font = Font(color=BLUE, underline="single", bold=True)
        cv[f"C{rr}"].value = d; cv.merge_cells(f"C{rr}:G{rr}")
    rr = r + len(sheets) + 3
    cv[f"B{rr}"].value = "COLOUR CODE"; cv[f"B{rr}"].font = Font(bold=True, color=BLUE)
    leg = [("Yellow fill", "Typed assumption (input)", None, "FFFF99"), ("Blue font", "Link pulled in from another row or sheet", "0000FF", None), ("Red font", "Key result row used by other sheets", "FF0000", None), ("Black font", "Intermediate calculation", "000000", None)]
    for i, (a, b, fc, bg) in enumerate(leg):
        cv[f"B{rr+1+i}"].value = a; cv[f"C{rr+1+i}"].value = b
        if fc: cv[f"B{rr+1+i}"].font = Font(color=fc, bold=True)
        if bg: cv[f"B{rr+1+i}"].fill = PatternFill("solid", fgColor=bg)
    rr += 7
    cv[f"B{rr}"].value = "HOW TO RUN"; cv[f"B{rr}"].font = Font(bold=True, color=BLUE)
    steps = ["1. Choose the scenario in InpC!F3 (drop-down). Everything recalculates.", "2. Change assumptions only in InpC (constants) and InpS (time series).",
             "3. Confirm the status line above says ALL CHECKS PASSED.", "4. To store a run: copy Track!H8:H69, paste values into M, N or O (Paste Special, Values)."]
    if kind == "pf": steps.append("5. After changing price, cost, tenor or DSCR inputs: Goal Seek Sculpt!F32 to 0 by changing the active scenario's gearing cell on InpC row 72 (J72 Base, K72 Downside, L72 Upside); then set the facility on InpC row 131 to ROUNDUP(Sculpt!F29, 0) for that scenario and repeat both until stable (the facility drives the fees, which drive the project cost).")
    for i, s in enumerate(steps): cv[f"B{rr+1+i}"].value = s; cv.merge_cells(f"B{rr+1+i}:G{rr+1+i}")
    rr += len(steps) + 2
    cv[f"B{rr}"].value = "CHANGE LOG (v1.0, 8 October 2026)"; cv[f"B{rr}"].font = Font(bold=True, color=BLUE)
    for i, ch in enumerate(changes):
        cv[f"B{rr+1+i}"].value = f"{i+1}. {ch}"; cv.merge_cells(f"B{rr+1+i}:G{rr+1+i}")
        cv[f"B{rr+1+i}"].alignment = Alignment(wrap_text=True, vertical="top"); cv.row_dimensions[rr + 1 + i].height = 45
    # ------------------------------------------------------------ Dashboard
    ds.column_dimensions["A"].width = 2; ds.column_dimensions["B"].width = 34
    for i in range(3, 25): ds.column_dimensions[L(i)].width = 11
    for rr_ in range(1, 4):
        for c in range(1, 25): ds.cell(rr_, c).fill = navy
    ds["B2"].value = "TrainStat: dashboard"; ds["B2"].font = Font(name="Georgia", size=20, bold=True, color="FFFFFF")
    ds["F2"].value = '="Scenario: "&InpC!F3'; ds["F2"].font = Font(size=13, bold=True, color="F0C06A")
    ds["K2"].value = "=Cover!C7"; ds["K2"].font = Font(size=12, bold=True, color="FFFFFF")
    kp = [("Shareholder NPV ($000s)", "=Val!F{}".format(VL(39) if kind == "pf" else 39), NUM), ("Shareholder IRR (XIRR)", "=Val!F{}".format(VL(49) if kind == "pf" else 49), "0.0%"),
          ("Minimum DSCR", "=Ratios!F13" if kind == "pf" else "=FinStat!F92", '0.00"x"'), ("Total CFADS ($000s)", "=Track!H10", NUM), ("Senior debt at close ($000s)", "=Debt!F11", NUM)]
    if kind == "pf": kp += [("Project cost at COD ($000s)", "=Constr!F52", NUM), ("Gearing at COD", "=Constr!F53/Constr!F52", "0.0%"), ("LLCR at COD", "=Ratios!F22", '0.00"x"')]
    # tiles in a grid: 4 per row
    for i, (lab, f, fmt) in enumerate(kp):
        row = 5 + (i // 4) * 4; c0 = 2 + (i % 4) * 5 if i % 4 else 2
        c0 = [2, 5, 10, 15][i % 4]
        span = 2 if c0 == 2 else 4
        for rr_ in (row, row + 1, row + 2):
            for cc in range(c0, c0 + span if c0 != 2 else 3):
                ds.cell(rr_, cc).fill = mist
        ds.cell(row, c0).value = lab; ds.cell(row, c0).font = Font(size=10, bold=True, color="5C6B7A")
        ds.cell(row + 1, c0).value = f; ds.cell(row + 1, c0).number_format = fmt; ds.cell(row + 1, c0).font = Font(name="Georgia", size=18, bold=True, color=NAVY)
    # annual table
    top = 14
    ds[f"B{top}"].value = "ANNUAL SUMMARY ($000s)"; ds[f"B{top}"].font = Font(bold=True, color=BLUE)
    hdr = top + 1
    ds[f"B{hdr}"].value = "Financial year"
    years = list(range(2022, 2042))
    for i, y in enumerate(years): ds.cell(hdr, 3 + i).value = y
    for c in range(2, 3 + len(years)):
        ds.cell(hdr, c).fill = navy; ds.cell(hdr, c).font = W; ds.cell(hdr, c).alignment = Alignment(horizontal="center")
    fsr = (lambda r: FS(r)) if kind == "pf" else (lambda r: r)
    rows = [("Operating revenue", f"FinStat!$J${fsr(9)}:$CO${fsr(9)}", 1), ("Operating cost", f"FinStat!$J${fsr(10)}:$CO${fsr(10)}", 1), ("EBITDA", f"FinStat!$J${fsr(11)}:$CO${fsr(11)}", 1),
            ("CFADS", f"FinStat!$J${fsr(38)}:$CO${fsr(38)}", 1), ("Senior debt service", f"FinStat!$J${fsr(88)}:$CO${fsr(88)}", 1), ("Dividends", f"FinStat!$J${fsr(47) if kind=='start' else FS(47)}:$CO${fsr(47) if kind=='start' else FS(47)}", -1)]
    for k, (lab, rng, sgn) in enumerate(rows):
        rr_ = hdr + 1 + k
        ds.cell(rr_, 2).value = lab
        for i, y in enumerate(years):
            cl = L(3 + i)
            ds.cell(rr_, 3 + i).value = f"={'-' if sgn<0 else ''}SUMIFS({rng},FinStat!$J$4:$CO$4,{cl}${hdr})"
            ds.cell(rr_, 3 + i).number_format = NUM
    rr_ = hdr + 1 + len(rows)
    ds.cell(rr_, 2).value = "DSCR (CFADS / debt service)"
    for i, y in enumerate(years):
        cl = L(3 + i)
        ds.cell(rr_, 3 + i).value = f"=IF(ABS({cl}{hdr+5})>CHK_TOL,{cl}{hdr+4}/{cl}{hdr+5},0)"; ds.cell(rr_, 3 + i).number_format = '0.00"x"'
    for rrr in range(hdr + 1, rr_ + 1):
        ds.cell(rrr, 2).font = Font(bold=True)
        for cc in range(2, 3 + len(years)):
            if (rrr - hdr) % 2 == 0: ds.cell(rrr, cc).fill = mist
    cats = Reference(ds, min_col=3, max_col=2 + len(years), min_row=hdr)
    def chart(kind_, title, rows_, anchor, ytitle="$000s", w=17, h=8.5):
        ch = BarChart() if kind_ == "bar" else LineChart()
        ch.title = title; ch.height = h; ch.width = w
        for rrr, color in rows_:
            ref = Reference(ds, min_col=2, max_col=2 + len(years), min_row=rrr)
            ch.add_data(ref, titles_from_data=True, from_rows=True)
            ch.series[-1].graphicalProperties.solidFill = color; ch.series[-1].graphicalProperties.line.solidFill = color
        ch.set_categories(cats); ch.y_axis.title = ytitle; ch.legend.position = "b"
        ds.add_chart(ch, anchor); return ch
    base = hdr + len(rows) + 4
    chart("bar", "Revenue and EBITDA by year", [(hdr + 1, "0B5E82"), (hdr + 3, "D99A2B")], f"B{base}")
    chart("bar", "CFADS and senior debt service", [(hdr + 4, "141B4D"), (hdr + 5, "D99A2B")], f"J{base}")
    chart("line", "DSCR by year", [(hdr + 7, "0B5E82")], f"B{base+18}", "x")
    chart("bar", "Dividends by year", [(hdr + 6, "1C7293")], f"J{base+18}")
    wb.move_sheet("Cover", offset=-wb.index(wb["Cover"]));
    return wb

def common_polish(wb):
    import re as _re
    for row in wb["Track"].iter_rows(min_row=1, max_row=80, min_col=10, max_col=12):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="): c.value = _re.sub(r">\s*0\.001", "> CHK_TOL", c.value)
    for ws in wb.worksheets:
        if ws.title in ("Cover", "Dashboard"): continue
        ws.sheet_view.showGridLines = False
        if ws.title in ("InpC", "InpS"): ws.sheet_properties.tabColor = "F2C97D"
        elif ws.title in ("Time", "Esc"): ws.sheet_properties.tabColor = TEAL
        elif ws.title in ("FinStat", "Val", "Track", "Ratios"): ws.sheet_properties.tabColor = NAVY
        elif ws.title != "Temp": ws.sheet_properties.tabColor = BLUE
    # hide the scratch sheet if it holds nothing useful
    if "Temp" in wb.sheetnames:
        t = wb["Temp"]; t.sheet_state = "hidden"

if __name__ == "__main__":
    kind, out = sys.argv[1], sys.argv[2]
    sol = json.load(open("/private/tmp/claude-501/s/model/pf2_solution.json"))
    wb = openpyxl.load_workbook("/private/tmp/claude-501/s/model/stage1.xlsx")
    changes = json.load(open("/private/tmp/claude-501/s/model/changes1.json"))
    scen = "Base: inflation 7% falling to 3%, fare US$7.15, O&M as contracted.  Downside: fare -10% and O&M +10%.  Upside: fare +10%."
    if kind == "pf":
        build_pf(wb, sol)
        track_pf(wb)
        changes = changes + ["PF extension added: Constr (construction period, fees, IDC, sources and uses), Sculpt (debt capacity, sizing, sculpted repayment), DSRA, SHL (shareholder loan), Ratios (DSCR, LLCR, PLCR); FinStat and Val extended; opening balances now flow from construction.",
                             "Debt sizing uses CFADS with tax before interest deduction, which avoids a circular reference; reported ratios use actual after-tax CFADS.",
                             "Facility rule: the facility on InpC row 131 is the commitment that can be drawn, ROUNDUP(MIN(debt capacity, gearing cap x project cost)); fees are charged on it. Previously the Upside facility was set to the DSCR capacity although the gearing cap binds, which overstated the Upside fees.",
                             "Dashboard: DSCR row now guards against float noise after maturity (=IF(ABS(debt service)>CHK_TOL, CFADS/debt service, 0)); it showed values of about -1e16 before.",
                             "Cover status now also tests Constr!F55 and the sources-less-uses row; Track flags use the named tolerance CHK_TOL instead of a typed 0.001.",
                             "Conventions retained and documented in the Prompt Book: all sponsor funding is dated at COD in Val (equity is actually paid in during construction, so the blended IRR is measured from COD); tax losses are refunded in cash (no loss carry-forward); the model has no DSRA draw and no dividend lock-up."]
    add_cover_dashboard(wb, kind, changes, scen)
    common_polish(wb)
    wb.save(out)
    print("saved", out)
