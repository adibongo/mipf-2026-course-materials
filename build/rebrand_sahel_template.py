# -*- coding: utf-8 -*-
"""Rebrand the Sahel Sun capital-stack template for MIPF Day 1 (formulas untouched)."""
import sys, openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Alignment
SRC, DST, LOGOS = sys.argv[1:4]
wb = openpyxl.load_workbook(SRC)
NAVY, TEAL, GOLD = "141B4D", "1C7293", "D99A2B"
CONTACT = "Prof Bongo Adi  |  +234 706 164 8699  |  badi@lbs.edu.ng  |  bongo.adi@fulbrightmail.org"
ws = wb["Read Me"]
ws.insert_rows(1, 3)                      # room for the logos; Read Me has no formulas
ws.row_dimensions[1].height = 54
for col, f, h in (("A", "lbs-logo.png", 48), ("B", "mipf_hex_256.png", 48), ("C", "bongo_ai_hex_256.png", 48)):
    im = Image(f"{LOGOS}/{f}"); r = h / im.height; im.height = h; im.width = int(im.width * r)
    ws.add_image(im, f"{col}1")
ws["A4"] = "Project Sahel Sun — 50MW Solar IPP: Capital Stack & Debt Sizing Template"
ws["A6"] = "MIPF Boot Camp, Day 1 — Infrastructure Finance: Structuring, Risk Management & Sustainable Project Delivery (14 October 2026)"
ws["A7"] = "Facilitator: Prof. Bongo Adi, Lagos Business School, Pan-Atlantic University"
ws["A8"] = None
# contact block under the existing text (last text row after the shift is 20)
last = ws.max_row + 2
ws.cell(last, 1, "Mastering Infrastructure and Project Finance (MIPF) Boot Camp  |  Lagos Business School  |  14 to 22 October 2026").font = Font(bold=True, color=NAVY)
ws.cell(last + 1, 1, CONTACT).font = Font(color=TEAL)
ws.cell(last + 2, 1, "Teaching case; all figures illustrative.").font = Font(italic=True, color="5C6B7A")
for s in wb:
    s.sheet_properties.tabColor = {"Read Me": GOLD, "Inputs": TEAL}.get(s.title, NAVY)
    s.oddHeader.left.text = "MIPF Boot Camp  |  Lagos Business School"
    s.oddHeader.right.text = "Day 1  |  Sahel Sun capital-stack template"
    s.oddFooter.center.text = CONTACT
    s.oddFooter.center.size = 8
wb.properties.title = "Project Sahel Sun: capital stack and debt sizing template (MIPF Boot Camp, Day 1)"
wb.properties.creator = "Prof Bongo Adi, Lagos Business School"
wb.calculation.fullCalcOnLoad = True
wb.active = 0
wb.save(DST)
print("saved", DST)
