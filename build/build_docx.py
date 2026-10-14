import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from mipf_data import *
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT=sys.argv[1]; M=OUT+"/src/media/"
NAVY=RGBColor(0x14,0x1B,0x4D); BLUE=RGBColor(0x0B,0x5E,0x82); GOLD=RGBColor(0xA8,0x70,0x0F); SLATE=RGBColor(0x5C,0x6B,0x7A)
d=Document()
sec=d.sections[0]; sec.different_first_page_header_footer=True; sec.page_width=Cm(21.0); sec.page_height=Cm(29.7)
for a in ("left_margin","right_margin","top_margin","bottom_margin"): setattr(sec,a,Inches(1))
st=d.styles["Normal"]; st.font.name="Garamond"; st.font.size=Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"),"Garamond")
st.paragraph_format.space_after=Pt(6)

def run(p,t,b=False,i=False,sz=None,c=None,caps=False):
    r=p.add_run(t); r.bold=b; r.italic=i
    r.font.name="Garamond"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"Garamond")
    if sz: r.font.size=Pt(sz)
    if c is not None: r.font.color.rgb=c
    if caps: r.font.small_caps=True
    return r
def para(t="",b=False,i=False,sz=None,c=None,al=None,after=None,before=None,keep=False):
    p=d.add_paragraph()
    if t: run(p,t,b,i,sz,c)
    if al=="c": p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if al=="j": p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None: p.paragraph_format.space_after=Pt(after)
    if before is not None: p.paragraph_format.space_before=Pt(before)
    if keep: p.paragraph_format.keep_with_next=True
    return p
def rule(p,color="D99A2B",sz=12):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement("w:pBdr"); e=OxmlElement("w:bottom")
    for k,v in (("val","single"),("sz",str(sz)),("space","1"),("color",color)): e.set(qn("w:"+k),v)
    b.append(e); pPr.append(b)
def head(t,sz=15):
    p=para(t,b=True,sz=sz,c=NAVY,before=14,after=6,keep=True); rule(p); return p
def pagebreak(): d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def pic(path,w,al="c"):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if al=="c" else WD_ALIGN_PARAGRAPH.LEFT
    p.add_run().add_picture(path,width=Inches(w)); return p

# footer / header
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
run(fp,"Lagos Business School  |  MIPF Boot Camp  |  14–22 October 2026  |  Prof Bongo Adi  +234 706 164 8699  |  badi@lbs.edu.ng  |  bongo.adi@fulbrightmail.org",sz=8,c=SLATE)
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
hp.add_run().add_picture(M+"mipf_hex_256.png",height=Cm(1.0)); hp.add_run("  "); hp.add_run().add_picture(M+"bongo_ai_hex_256.png",height=Cm(1.0)); hp.add_run("  "); hp.add_run().add_picture(M+"lbs-logo.png",height=Cm(1.1))

# ------------------------------------------------ cover
pic(M+"lbs-logo.png",3.1,"l"); para(after=14)
tbh=d.add_table(rows=1,cols=2); tbh.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,im in enumerate(["mipf_hex.png","bongo_ai_hex.png"]):
    c=tbh.rows[0].cells[i]; c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraphs[0].add_run().add_picture(M+im,width=Inches(1.5))

t=para(TITLE,b=True,sz=26,c=NAVY,al="c",before=6,after=4)
para(VENUE,sz=14,c=BLUE,al="c",after=0); para(PLACE,sz=13,c=BLUE,al="c",after=10)
p=para(DATES_LONG,b=True,sz=16,c=GOLD,al="c",after=16); 
tb=d.add_table(rows=1,cols=3); tb.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,im in enumerate(["image1.png","image2.png","image3.png"]):
    c=tb.rows[0].cells[i]; c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    c.paragraphs[0].add_run().add_picture(M+im,width=Inches(1.85))
para(after=30)
p=para("Course information and syllabus",b=True,sz=20,c=NAVY,al="c",after=4); rule(p)
para("Open-enrolment programme  |  Seven teaching days: Wednesday 14th to Friday 16th and Monday 19th to Thursday 22nd October 2026 (Saturday 17th and Sunday 18th are rest days)",i=True,sz=11,c=SLATE,al="c")
pagebreak()

# ------------------------------------------------ narrative (2024/2025 order)
head("Introduction")
for t in INTRO: para(t,al="j")
head("Objective"); para(OBJECTIVE,al="j")
head("Overview")
for t in OVERVIEW: para(t,al="j")
para("Modules:",b=True,c=NAVY,keep=True)
for t in MODULES_INTRO: para(t,al="j")
for h,t in MODULES:
    para(h+":",b=True,c=BLUE,before=6,after=2,keep=True); para(t,al="j")
para(CASES_LINE,al="j")
para("Capstone Project:",b=True,c=BLUE,before=6,after=2,keep=True); para(CAPSTONE,al="j")
para(CLOSING,al="j")
para("Prof Bongo Adi",b=True,after=0,before=10); para("Course Director",i=True)
head("Contact"); 
for t in ("Prof Bongo Adi, Course Director, Lagos Business School, Pan-Atlantic University","Phone / WhatsApp: +234 706 164 8699","Email: badi@lbs.edu.ng","Fulbright email: bongo.adi@fulbrightmail.org"): para(t,after=2)
head("Participant Profile"); para(PROFILE,al="j")

# at a glance table
head("The week at a glance")
tbl=d.add_table(rows=1,cols=4); tbl.style="Table Grid"
for i,h in enumerate(["Day","Date","Module / focus","Facilitator"]):
    c=tbl.rows[0].cells[i]; c.text=""; run(c.paragraphs[0],h,b=True,c=RGBColor(255,255,255),sz=10)
    sh=OxmlElement("w:shd"); sh.set(qn("w:val"),"clear"); sh.set(qn("w:fill"),"141B4D"); c._tc.get_or_add_tcPr().append(sh)
for dy in DAYS:
    r=tbl.add_row().cells
    for i,v in enumerate([str(dy["n"]),dy["date"],dy["mtitle"],dy["fac"]]):
        r[i].text=""; run(r[i].paragraphs[0],v,sz=10)
for r in tbl.rows:
    for i,w in enumerate([0.5,1.9,3.0,1.6]): r.cells[i].width=Inches(w)
para(SCHEDULE_NOTE,i=True,sz=9.5,c=SLATE,before=6)

# faculty
pagebreak()
head("Faculty")
photos={"Prof. Bongo Adi":"image4.png","Promise Godbold":"image5.jpeg","Michelle Chikezie, LLB, LLM":"chikezie.jpg"}
for f in FACULTY:
    para(f["name"]+" – "+f["role"],b=True,sz=12.5,c=NAVY,before=10,after=4,keep=True)
    if f["name"] in photos: pic(M+photos[f["name"]],1.25,"l")
    for t in f["bio"]: para(t,al="j",i=t.startswith("["))

# ------------------------------------------------ schedule
def ampm(s,e):
    h=int(e.split(".")[0]); return "pm" if (h>=12 or h<=7) else "am"
def trange(s,e): return f"{s} – {e} {ampm(s,e)}"
def label(keys,pref):
    out=[]
    for k in keys:
        r=R[k]
        if r["kind"]=="web": out.append(f'{r["t"]} ({r["s"]}): {r["u"]}')
        elif r["kind"]=="hosted": out.append(f'{r["t"]} (on the course website)')
        elif r.get("od"): out.append(f'{r["t"]} (OneDrive: ' + "; ".join(u for l,u in r["od"]) + ')')
        else: out.append(f'{r["t"]} (Harvard case, in class)')
    return out
for dy in DAYS:
    pagebreak()
    para("Mastering Infrastructure and Project Finance",b=True,sz=13,c=BLUE,al="c",after=0)
    p=para(dy["mtitle"],b=True,sz=17,c=NAVY,al="c",after=2); 
    p=para(f'Day {dy["n"]}  |  {dy["date"]}',b=True,sz=13,c=GOLD,al="c",after=6); rule(p)
    para(dy["blurb"],i=True,sz=10.5,c=SLATE,al="j",after=8)
    for s in dy["sessions"]:
        p=d.add_paragraph(); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.left_indent=Inches(1.75); p.paragraph_format.first_line_indent=Inches(-1.75)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(1.75))
        run(p,trange(s["start"],s["end"]),b=True,c=NAVY); run(p,"\t"); run(p,s["title"],b=(s["kind"]!="break"),c=(NAVY if s["kind"]!="break" else SLATE))
        if s["kind"]=="break": continue
        def line(lbl,txt,it=False):
            q=d.add_paragraph(); q.paragraph_format.left_indent=Inches(1.75); q.paragraph_format.space_after=Pt(1)
            run(q,lbl,b=True,sz=10.5,c=BLUE); run(q,txt,i=it,sz=10.5); return q
        line("In class:  ",s["mode"]); line("Facilitator:  ",s["fac"])
        if s["desc"]: line("Focus:  ",s["desc"])
        if s["exercise"]: line("Exercise:  ",s["exercise"])
        if s["sheets"]: line("TrainStat sheets:  ",", ".join(s["sheets"]))
        if s["note"]: line("Note:  ",s["note"],True)
        if s["read"]:
            q=line("Reading materials:  ",""); 
            for j,x in enumerate(label(s["read"],"R"),1):
                z=d.add_paragraph(); z.paragraph_format.left_indent=Inches(2.0); z.paragraph_format.space_after=Pt(0); run(z,f"{j}. {x}",sz=10)
        if s["alt"]:
            q=line("Alternative reading (cases):  ","")
            for j,x in enumerate(label(s["alt"],"A"),1):
                z=d.add_paragraph(); z.paragraph_format.left_indent=Inches(2.0); z.paragraph_format.space_after=Pt(0); run(z,f"{j}. {x}",sz=10)

# ------------------------------------------------ appendices
pagebreak(); head("Appendix A: The TrainStat start model")
para("The modelling days (3 to 7) use one workbook, TrainStat_StartModel.xlsx, which participants receive before Day 3. Days 4 and 5 complete the operating three-statement model; Days 6 and 7 extend it into a project finance model.",al="j")
for k,v in MODEL["facts"]:
    q=para(after=1); run(q,k+":  ",b=True,c=BLUE,sz=10.5); run(q,v,sz=10.5)
para("Workbook map",b=True,c=NAVY,before=10,keep=True)
tb2=d.add_table(rows=1,cols=3); tb2.style="Table Grid"
for i,h in enumerate(["Sheet","Purpose","Contents"]):
    c=tb2.rows[0].cells[i]; c.text=""; run(c.paragraphs[0],h,b=True,sz=10,c=RGBColor(255,255,255))
    sh=OxmlElement("w:shd"); sh.set(qn("w:val"),"clear"); sh.set(qn("w:fill"),"141B4D"); c._tc.get_or_add_tcPr().append(sh)
for a,b_,c_ in MODEL["sheets"]:
    r=tb2.add_row().cells
    for i,v in enumerate([a,b_,c_]): r[i].text=""; run(r[i].paragraphs[0],v,sz=10)

head("Appendix B: Reading list")
para("Course pack items are distributed in class and are not posted on the website. Public links were checked at the time of writing.",i=True,sz=10,c=SLATE)
for kind,ttl in (("hosted","Provided on the course website"),("pack","Course pack and OneDrive readings"),("web","Public links")):
    para(ttl,b=True,c=BLUE,before=8,after=2,keep=True)
    for k,r in R.items():
        if r["kind"]==kind:
            q=para(after=1); run(q,"•  "+r["t"],sz=10.5); run(q,"  – "+r["s"],i=True,sz=10,c=SLATE)
            if kind=="web": run(q,"  "+r["u"],sz=9.5,c=BLUE)
out=OUT+"/MIPF_October_2026_Course_Syllabus.docx"; d.save(out); print("saved",out)
