# -*- coding: utf-8 -*-
"""Add the LBS logo to the header and the contact line to the footer of the Day 1 Word documents."""
import sys, copy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
SRC, DST, LOGO = sys.argv[1:4]
d = Document(SRC)
seen = set()
for sec in d.sections:
    if id(sec.header._element) in seen or sec.header.is_linked_to_previous: continue
    seen.add(id(sec.header._element))
    hp = sec.header.paragraphs[0]
    run = hp.add_run()
    run.add_picture(LOGO, height=Cm(1.0))
    hp._p.remove(run._r)
    ppr = hp._p.pPr
    hp._p.insert(list(hp._p).index(ppr) + 1 if ppr is not None else 0, run._r)   # LBS logo first, hexes after
    hp.add_run("  ")
    fp = sec.footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Prof Bongo Adi  |  +234 706 164 8699  |  badi@lbs.edu.ng  |  bongo.adi@fulbrightmail.org")
    r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x1C, 0x72, 0x93)
d.save(DST)
print("saved", DST)
