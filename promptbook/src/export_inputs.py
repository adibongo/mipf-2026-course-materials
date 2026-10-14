"""Export model inputs (3 scenarios) and the formula atlas from the verified workbooks to JSON for the engine, app and prompt book."""
import openpyxl, json, datetime as dt, re, sys, warnings
warnings.filterwarnings("ignore")
M = "/private/tmp/claude-501/s/model"
OUT = sys.argv[1]
wb = openpyxl.load_workbook(f"{M}/pf_cached.xlsx", data_only=True)
c, s = wb["InpC"], wb["InpS"]
iso = lambda v: v.date().isoformat() if isinstance(v, (dt.datetime, dt.date)) else v
ROWS = {  # key: InpC row
 "first": 9, "acq": 10, "nyrs": 11, "mpp": 12, "cpi0": 23, "cpiH": 25, "cpiO": 27, "capD": 31, "capI": 33, "fare": 35, "arDays": 37,
 "hFixed": 42, "hVar": 44, "omSens": 46, "omFixed": 48, "apDays": 50, "nca0": 64, "life": 65, "lifeCapex": 67, "gearing": 72, "profile": 74,
 "rate": 77, "dc": 79, "tdep": 84, "tax": 86, "valDate": 90, "yDays": 91, "req": 93,
 "constrN": 113, "hard": 115, "upfront": 127, "commit": 129, "facility": 131, "drawdown": 133, "eqCommit": 135, "tenor": 139, "dscr": 140,
 "gmax": 141, "dsraN": 143, "shlPct": 147, "shlRate": 148}
P = {}
for scn, col in (("Base", 10), ("Downside", 11), ("Upside", 12)):
    d = {k: iso(c.cell(r, col).value) for k, r in ROWS.items()}
    d["infl"] = [[iso(c.cell(r, 5).value), c.cell(r, col).value] for r in range(17, 22)]
    d["capex"] = [[iso(c.cell(r, 5).value), c.cell(r, col).value] for r in range(56, 63)]
    d["spend"] = [c.cell(r, col).value for r in range(118, 126)]
    P[scn] = d
S = {"dates": [iso(s.cell(2, 10 + i).value) for i in range(84)]}
for k, r in (("td", 8), ("ti", 9), ("ad", 11), ("ai", 12), ("od", 14), ("oi", 15), ("p1", 18), ("p2", 19), ("p3", 20)):
    S[k] = [s.cell(r, 10 + i).value or 0 for i in range(84)]
LAB = {k: str(c.cell(r, 5).value) for k, r in ROWS.items()}
UNIT = {k: str(c.cell(r, 7).value or "") for k, r in ROWS.items()}
import re as _re
_funcs = set()
_wf = openpyxl.load_workbook(f"{M}/pf_formulas.xlsx")
for _ws in _wf.worksheets:
    for _row in _ws.iter_rows(min_row=1, max_row=min(_ws.max_row, 160), max_col=20):
        for _c in _row:
            if isinstance(_c.value, str) and _c.value.startswith("="): _funcs.update(_re.findall(r"([A-Z][A-Z0-9.]+)\(", _c.value.replace("_xlfn.", "")))
_funcs -= {"CELL", "SEARCH", "LEN", "RIGHT", "LEFT", "NOW", "TODAY", "MID"}
json.dump(dict(P=P, S=S, labels=LAB, units=UNIT, rows=ROWS, funcs=sorted(_funcs) + ["FORMULATEXT"]), open(f"{OUT}/inputs.json", "w"), separators=(",", ":"), default=str)
def cl(f):
    if f is None: return ""
    f = str(f).strip()
    if f.startswith("{") and f.endswith("}"): f = f[1:-1]
    out, q = [], False
    for ch in f:
        if ch == '"': q = not q
        if ch == " " and not q: continue
        out.append(ch)
    return "".join(out).replace("_xlfn.", "")
def num(v):
    try: return round(float(v), 4)
    except Exception: return v
def colnum(c):
    n = 0
    for ch in c: n = n * 26 + ord(ch) - 64
    return n
REF = re.compile(r"(?<![A-Za-z_])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(A-Za-z_])")
def rc(f, col):
    """Relative form of a formula as seen from column number `col` (row-relative refs left as-is: rows do not shift along a row)."""
    def sub(m):
        ca, cl_, ra, rw = m.groups(); n = colnum(cl_)
        return f"{'C'+str(n) if ca else 'C['+str(n-col)+']'}{ra}{rw}"
    return REF.sub(sub, f)
def kind(lab, d, f, j):
    if not f and not j and str(lab).isupper() and len(str(lab)) > 3: return "section"
    if j and re.fullmatch(r"=[A-Za-z]*!?\$?[A-Z]{1,3}\$?\d+", j): return "pull"
    if j: return "calc"
    if f: return "const"
    return "label"
def make_atlas(rows):
    atlas = []
    for sh, rr in rows.items():
        for r, d in rr.items():
            lab = d.get("E_v") or d.get("B") or d.get("A")
            f = {k: cl(d.get(k)) for k in ("F", "J", "K", "L") if d.get(k) not in (None, "")}
            if not lab and not f: continue
            if str(lab).startswith("=") and "CELL(" in str(lab): continue
            item = dict(s=sh, r=int(r), l=str(lab or ""), f=f.get("F", ""), j=f.get("J", ""), t=kind(lab, d, f.get("F", ""), f.get("J", "")))
            if f.get("J") and f.get("K") and rc(f["J"], 10) != rc(f["K"], 11): item["x"] = f["K"]
            if f.get("K") and f.get("L") and rc(f["K"], 11) != rc(f["L"], 12): item["x2"] = f["L"]
            if "H_v" in d and d["H_v"] not in (None, "", "empty"): item["h"] = num(d["H_v"])
            if "F_v" in d and d["F_v"] not in (None, "", "empty", "Constant"): item["fv"] = num(d["F_v"])
            if "L_v" in d: item["lv"] = num(d["L_v"])
            for kk, key in (("S_v", "v23"), ("AY_v", "v31"), ("CO_v", "vco")):
                if kk in d and isinstance(d[kk], (int, float)) and not isinstance(d[kk], bool): item[key] = num(d[kk])
            u = d.get("G_v")
            if isinstance(u, str) and u and not u.startswith("=") and u not in ("unit",): item["u"] = u
            if sh == "InpC" and any(d.get(k) not in (None, "") for k in "JKL"):
                item["sc"] = [str(d.get(k))[:19] if d.get(k) not in (None, "") else "" for k in "JKL"]
            atlas.append(item)
    return atlas
A_pf = make_atlas(json.load(open(f"{M}/pf_rows.json")))
A_st = make_atlas(json.load(open(f"{M}/start_rows.json")))
json.dump(A_pf, open(f"{OUT}/atlas_pf.json", "w"), separators=(",", ":"))
json.dump(A_st, open(f"{OUT}/atlas_start.json", "w"), separators=(",", ":"))
print("inputs ok; atlas rows pf", len(A_pf), "start", len(A_st))
